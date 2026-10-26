"""Validation tests for the flat skill-level parameter authoring table."""

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from tools.skill_level_parameters import materialize_nested_workbook


class SkillLevelParameterTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.flat = self.root / "flat.xlsx"
        self.nested = self.root / "nested.xlsx"
        self.write_flat([
            (1, 10, 1, 0, "Damage", 12),
            (2, 10, 1, 1, "Radius", 2),
            (3, 10, 2, 0, "Damage", 18),
            (4, 10, 2, 1, "Radius", 3),
        ])
        self.write_nested()

    def tearDown(self):
        self.temp_dir.cleanup()

    def write_flat(self, data):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Id", "SkillId", "Level", "ParamIndex", "ParamName", "Value"])
        sheet.append(["int", "int", "int", "int", "string", "float"])
        sheet.append(["", "", "", "", "", ""])
        for row in data:
            sheet.append(row)
        workbook.save(self.flat)
        workbook.close()

    def write_nested(self, levels=2):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([None, None, None, None, "ParamNames"])
        sheet.append([None] * 6)
        sheet.append([None] * 6)
        for level in range(1, levels + 1):
            row = 4 + level
            sheet.cell(row, 2).value = 10 if level == 1 else None
            sheet.cell(row, 3).value = 0
            sheet.cell(row, 4).value = 0
        workbook.save(self.nested)
        workbook.close()

    def materialize_rows(self, rows):
        self.write_flat(rows)
        materialize_nested_workbook(self.flat, self.nested)

    def test_materializes_flat_parameters_into_nested_values(self):
        materialize_nested_workbook(self.flat, self.nested)
        workbook = load_workbook(self.nested, data_only=True)
        try:
            sheet = workbook.active
            self.assertEqual(10, sheet.cell(5, 2).value)
            self.assertEqual("12.0,2.0", sheet.cell(5, 5).value)
            self.assertEqual("Damage,Radius", sheet.cell(5, 6).value)
            self.assertIsNone(sheet.cell(6, 2).value)
            self.assertEqual("18.0,3.0", sheet.cell(6, 5).value)
        finally:
            workbook.close()

    def test_rejects_duplicate_row_id_and_composite_key(self):
        with self.assertRaisesRegex(ValueError, "duplicate row Id"):
            self.materialize_rows([(1, 10, 1, 0, "Damage", 12), (1, 10, 1, 1, "Radius", 2)])
        with self.assertRaisesRegex(ValueError, "duplicate skill parameter row"):
            self.materialize_rows([(1, 10, 1, 0, "Damage", 12), (2, 10, 1, 0, "Damage", 13)])

    def test_rejects_non_contiguous_indexes_and_inconsistent_names(self):
        with self.assertRaisesRegex(ValueError, "contiguous from zero"):
            self.materialize_rows([(1, 10, 1, 1, "Damage", 12), (2, 10, 2, 1, "Damage", 13)])
        with self.assertRaisesRegex(ValueError, "multiple names"):
            self.materialize_rows([(1, 10, 1, 0, "Damage", 12), (2, 10, 2, 0, "Power", 13)])
        with self.assertRaisesRegex(ValueError, "multiple indexes"):
            self.materialize_rows([(1, 10, 1, 0, "Damage", 12), (2, 10, 1, 1, "Damage", 13)])

    def test_rejects_incomplete_parameter_level_and_missing_metadata_level(self):
        with self.assertRaisesRegex(ValueError, "parameter set is incomplete"):
            self.materialize_rows([(1, 10, 1, 0, "Damage", 12), (2, 10, 1, 1, "Radius", 2),
                                   (3, 10, 2, 0, "Damage", 18)])
        self.write_flat([
            (1, 10, 1, 0, "Damage", 12),
            (2, 10, 1, 1, "Radius", 2),
            (3, 10, 2, 0, "Damage", 18),
            (4, 10, 2, 1, "Radius", 3),
        ])
        self.write_nested(levels=3)
        with self.assertRaisesRegex(ValueError, "do not match skill level metadata"):
            materialize_nested_workbook(self.flat, self.nested)

    def test_rejects_non_contiguous_levels(self):
        with self.assertRaisesRegex(ValueError, "levels must be contiguous"):
            self.materialize_rows([(1, 10, 1, 0, "Damage", 12), (2, 10, 3, 0, "Damage", 18)])


if __name__ == "__main__":
    unittest.main()
