"""Maintain Chinese field notes in the production Luban workbooks."""

import argparse
import re
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill


DATAS = Path(__file__).resolve().parents[1] / "Production/Datas"
NOTE_FILL = PatternFill("solid", fgColor="E8F1F2")

# These names carry semantics that cannot be inferred from their English tokens.
EXACT = {
    "Id": "本表记录的唯一 ID；其他表按该值引用。",
    "RowId": "导表行 ID；对象根或数组根重建时使用，不是玩法 ID。",
    "Code": "旧资源主键；保留原运行时编号。",
    "Name": "编辑器和调试时显示的名称。",
    "Key": "稳定的程序标识或业务键；修改前检查代码与引用。",
    "OmittedFields": "迁移兼容字段：逗号分隔的原 JSON 缺失字段名；新数据通常留空。",
    "NullFields": "迁移兼容字段：逗号分隔的原 JSON 显式 null 字段名；新数据通常留空。",
    "CastFlowId": "必填；施法 Pipeline ID，引用 Pipeline 编辑器发布的 skill_flows。",
    "PreCastFlowId": "可选；前置施法 Pipeline ID，0 表示不使用。",
    "LevelTableId": "技能等级表 ID；0 表示不使用等级参数。",
    "RequiredTargetQueryId": "施法所需目标查询模板 ID；0 表示无此要求。",
    "SkillButtonTemplateId": "技能按钮与瞄准表现模板 ID；0 表示无专用模板。",
    "TriggerIds": "触发器 ID 列表；由 Trigger Authoring 发布，逗号分隔。",
    "OnAddEffects": "Buff 添加时执行的触发器 ID 列表；历史字段名为 Effects。",
    "OnRemoveEffects": "Buff 移除时执行的触发器 ID 列表；历史字段名为 Effects。",
    "OnIntervalEffects": "Buff 周期触发器 ID 列表；历史字段名为 Effects。",
    "OnHitEffectId": "投射物命中时执行的触发器 ID；历史字段名为 Effect。",
    "IntervalTriggerIds": "持续流程每次间隔执行的触发器 ID 列表。",
    "OnDelayTriggerIds": "区域延迟结束时执行的触发器 ID 列表。",
    "OnEnterTriggerIds": "目标进入区域时执行的触发器 ID 列表。",
    "OnExitTriggerIds": "目标离开时执行的触发器 ID 列表。",
    "OnIntervalTriggerIds": "区域每次间隔执行的触发器 ID 列表。",
    "TagIds": "授予或携带的标签定义 ID 列表，逗号分隔。",
    "Tags": "可选标签 ID 列表；技能表仅按需填写 SkillSelection，召唤物和玩法表留空。",
    "PassiveSkillIds": "被动技能 ID 列表，引用被动技能表。",
    "SkillIds": "主动技能 ID 列表，引用技能表。",
    "ActiveSkills": "旧属性模板中的主动技能 ID；由角色表技能列表补齐。",
    "PassiveSkills": "旧属性模板中的被动技能 ID；由角色表技能列表补齐。",
    "ModifierIds": "有序修饰器定义 ID 列表；顺序影响应用次序。",
    "AttrScales": "召唤物继承属性的配置列表；每项指定属性 ID 与缩放值。",
    "ParamNames": "各等级 Params 的参数名，逗号分隔；名称及顺序必须与每级数值一一对应。",
    "Params": "本级技能参数数值列表；顺序与 ParamNames 严格一致。",
    "Levels": "技能逐级数值列表；每一行对应一个技能等级。",
    "ContinuousTagTemplateId": "持续标签模板 ID；0 表示不启用。",
    "PresentationTemplateId": "表现模板 ID；用于复用挂点、跟随与外观配置。",
    "AttributeTemplateId": "属性模板 ID，引用角色表中的属性模板。",
    "DefaultComponentTemplateIds": "默认组件模板 ID 列表，逗号分隔。",
    "ModelId": "模型资源 ID，引用模型表。",
    "VfxId": "特效资源 ID，引用特效资源表。",
    "AssetId": "特效资源 ID，引用特效资源表。",
    "TemplateId": "模板 ID；按本表用途引用对应模板表。",
    "SummonId": "召唤物 ID，引用召唤物表。",
    "StateMachineProfileId": "投射物状态机配置 ID；由状态机编辑器发布。",
    "RootNodeId": "本规则组的唯一根节点 ID。",
    "GroupId": "所属标签规则组 ID。",
    "ParentId": "父规则节点 ID；根节点填 0。",
    "RefGroupId": "GroupRef 节点引用的其他规则组 ID。",
    "TagId": "Tag 节点比较的标签定义 ID。",
    "RuleGroupId": "结果标志映射引用的标签规则组 ID。",
    "FlagId": "结果标志 ID，引用结果标志定义。",
    "BitIndex": "结果位在位掩码中的序号；发布后不可重排。",
    "Operator": "规则节点操作：Tag、All、Any、Not 或 GroupRef。",
    "Exact": "Tag 匹配方式：true 为精确匹配，false 允许层级匹配。",
    "Namespace": "标签命名空间：旧数字、游戏状态或技能选择。",
    "ValueKind": "属性值类型：Flat 为普通数值，Percent 为比率。",
    "DefaultValue": "属性未被模板覆盖时使用的默认数值。",
    "PrefabPath": "模型 Prefab 的资源路径。",
    "Resource": "特效资源路径，由表现层加载。",
    "Pointer": "资源节点在 JSON 中的 RFC 6901 指针。",
    "ValueJson": "资源节点的 JSON 标量值。",
    "Path": "输出资源文档路径。",
    "Kind": "当前记录的类型；取值由对应配置契约定义。",
    "Category": "技能分类编号；用于技能分组。",
    "SkillType": "技能类型枚举；按 strong_types.xml 中 SkillType 填写。",
    "Priority": "应用优先级；数值关系由消费系统决定。",
    "DurationMs": "持续时间，单位毫秒；0 的含义以对应运行时逻辑为准。",
    "CooldownMs": "冷却时间，单位毫秒。",
    "IntervalMs": "周期执行间隔，单位毫秒。",
    "DelayMs": "开始前延迟，单位毫秒。",
    "LifetimeMs": "实例存活时间，单位毫秒。",
    "DefaultDurationMs": "默认持续时间，单位毫秒。",
    "Range": "技能有效距离，使用战斗世界单位。",
    "Radius": "半径，使用战斗世界单位。",
    "Speed": "移动速度，使用战斗世界单位/秒。",
    "Scale": "表现缩放倍数；1 表示原始大小。",
    "OffsetX": "相对挂点或中心的 X 轴偏移。",
    "OffsetY": "相对挂点或中心的 Y 轴偏移。",
    "OffsetZ": "相对挂点或中心的 Z 轴偏移。",
    "ColorR": "特效颜色红色通道，范围 0～1。",
    "ColorG": "特效颜色绿色通道，范围 0～1。",
    "ColorB": "特效颜色蓝色通道，范围 0～1。",
    "ColorA": "特效颜色透明度通道，范围 0～1。",
    "Follow": "是否随挂接目标持续移动。",
    "Socket": "表现挂点名称；空值使用默认位置。",
    "BitIndex": "结果位索引；发布后必须保持稳定。",
}

TOKENS = {
    "Activation": "激活", "Active": "主动", "Aim": "瞄准", "Alive": "存活", "Angle": "角度",
    "Application": "应用", "Arc": "弧线", "Area": "区域", "Attack": "攻击", "Attr": "属性",
    "Attribute": "属性", "Base": "基础", "Basic": "普攻", "Battle": "战斗", "Blocked": "阻断",
    "Blood": "生命", "Bloodsucking": "吸血", "Bool": "布尔", "Bounds": "边界", "Brain": "AI 脑",
    "By": "依据", "Caster": "施法者", "Center": "中心", "Client": "客户端", "Close": "关闭",
    "Coefficient": "系数", "Collision": "碰撞", "Cols": "列数", "Component": "组件", "Config": "配置", "Context": "上下文",
    "Continuous": "持续", "Cooldown": "冷却", "Cost": "消耗", "Count": "数量", "Critical": "暴击",
    "Curve": "曲线", "Decay": "衰减", "Decision": "决策", "Defense": "防御", "Deg": "度",
    "Despawn": "销毁", "Dispel": "驱散", "Distance": "距离", "Drag": "拖拽", "Driver": "驱动器",
    "Duration": "持续时间", "Effect": "效果", "Emitter": "发射器", "Enable": "启用", "Evaluation": "求值", "Execution": "执行",
    "Exit": "退出", "Explicit": "显式", "Extra": "额外", "Face": "朝向", "Fan": "扇形",
    "Flag": "标志", "Float": "浮点", "Forward": "前方", "Frames": "帧", "Generate": "生成",
    "Grid": "网格", "Group": "组", "Half": "半", "Height": "高度", "Hero": "英雄",
    "Hit": "命中", "Host": "主机", "Hp": "生命值", "Icon": "图标", "Id": "ID",
    "Inherit": "继承", "Input": "输入", "Int": "整数", "Interval": "间隔", "Layer": "层", "Length": "长度",
    "Level": "等级", "Long": "长", "Magic": "法术", "Magnitude": "幅值", "Main": "主",
    "Mana": "法力", "Mask": "掩码", "Max": "最大", "Min": "最小", "Mode": "模式",
    "Model": "模型", "Motion": "运动", "Move": "移动", "Ms": "毫秒", "Network": "网络",
    "Object": "对象", "Of": "的", "On": "在", "Ongoing": "持续中", "Op": "操作",
    "Ops": "操作列表", "Order": "顺序", "Out": "脱离", "Overflow": "超限", "Owner": "拥有者",
    "Param": "参数", "Pattern": "模式", "Per": "每", "Percent": "百分比", "Penetration": "穿透",
    "Physics": "物理", "Player": "玩家", "Point": "点位", "Policy": "策略", "Port": "端口",
    "Position": "位置", "Prediction": "预测", "Press": "按压", "Priority": "优先级", "Process": "流程",
    "Profile": "配置", "Projectile": "投射物", "Provider": "提供器", "Query": "查询", "R": "比率",
    "Random": "随机", "Rate": "速率", "Recording": "录制", "Reduction": "减免", "Reduce": "缩减",
    "Ref": "引用", "Refresh": "刷新", "Removal": "移除", "Remove": "移除", "Require": "要求",
    "Required": "必需", "Response": "响应", "Resilience": "韧性", "Return": "返回", "Replay": "回放",
    "Rotation": "旋转", "Rows": "行数", "Rule": "规则", "Run": "运行", "Scales": "缩放配置",
    "Schema": "结构", "Score": "评分", "Scorers": "评分器列表", "Seconds": "秒", "Seed": "种子",
    "Select": "选择", "Selection": "选择", "Selector": "选择器", "Self": "自身", "Shape": "形状",
    "Shot": "发射", "Skill": "技能", "Source": "来源", "Spawn": "生成", "Spacing": "间距",
    "Stack": "叠层", "Stacking": "叠层", "Start": "开始", "State": "状态", "Stats": "属性",
    "Stop": "停止", "Sub": "子", "Summon": "召唤物", "Suppressed": "被抑制", "Sync": "同步",
    "Tag": "标签", "Target": "目标", "Team": "队伍", "Template": "模板", "Threshold": "阈值",
    "Tick": "逻辑帧", "Time": "时间", "To": "到", "Total": "总", "Trigger": "触发器",
    "Type": "类型", "Unit": "单位", "Upgrade": "升级", "Use": "使用", "Valid": "有效",
    "Value": "值", "Variant": "变体", "Variants": "变体", "Version": "版本", "Vfx": "特效",
    "View": "视图", "Walkable": "可行走", "Width": "宽度", "Win": "胜利", "World": "世界",
    "Yaw": "偏航角", "X": "X", "Y": "Y", "Z": "Z",
    "id": "ID", "count": "数量", "desc": "描述", "name": "名称", "schema": "结构",
    "states": "状态定义", "variants": "动作变体",
}

TOKENS.update({
    "After": "之后", "Areas": "区域", "At": "在", "Attach": "挂接", "Blocks": "阻挡", "Buff": "增益",
    "Candidate": "候选", "Category": "类别", "Code": "编号", "Combat": "战斗", "Condition": "条件",
    "Constant": "常量", "Debug": "调试", "Default": "默认", "Degrees": "角度", "Delay": "延迟",
    "Die": "死亡", "Display": "显示", "Domain": "领域", "Effective": "生效", "Emit": "发射", "Euler": "欧拉角",
    "Expire": "过期", "Grant": "授予", "Hits": "命中次数", "Ids": "ID 列表",
    "Indicator": "指示器", "Index": "索引", "Key": "键", "Kind": "类型",
    "Formula": "公式", "Launcher": "发射器", "Movement": "移动", "Name": "名称", "Objects": "对象",
    "Offset": "偏移", "Players": "玩家", "Points": "点位", "Radius": "半径", "Reference": "引用",
    "Resolve": "解析", "Role": "角色", "Rules": "规则", "Range": "范围", "Remaining": "剩余",
    "Second": "秒", "Size": "大小", "Slot": "槽位", "Speed": "速度", "Stacks": "层数",
    "Status": "状态", "Targets": "目标数", "Terms": "项", "Writable": "可写",
})

SPECIAL = {
    "full_name": "完整表或类型名，包含命名空间。", "parent": "父类型名；无继承时留空。",
    "valueType": "类型对应的值类型。", "sep": "复合值的分隔符。", "alias": "类型别名。",
    "comment": "供 Luban 读取的类型说明。", "flags": "是否按位组合的枚举。",
    "unique": "枚举数值是否必须唯一。", "group": "枚举所属分组。", "tags": "类型附加标签。",
    "value_type": "表记录的值类型。", "read_schema_from_file": "是否从输入工作簿读取结构定义。",
    "input": "表数据来源工作簿及子 sheet。", "index": "索引字段；留空时使用值类型首字段。",
    "mode": "表装载模式：one、map 或 list。", "output": "导出文件名或路径。",
    "*fields": "bean 的字段定义列表；子列分别填写字段名、类型、别名和注释。",
    "*items": "枚举成员列表；子列分别填写成员名、数值和注释。",
}


def describe(name, kind, sheet):
    if name in EXACT:
        return EXACT[name]
    if name in SPECIAL:
        return SPECIAL[name]
    if name.endswith("TagIds"):
        return f"{translate(name[:-6])}标签定义 ID 列表，逗号分隔。"
    if name.endswith("RuleGroupId"):
        return f"{translate(name[:-11])}条件规则组 ID；0 表示不使用。"
    if name.endswith("TriggerIds"):
        return f"{translate(name[:-10])}触发器 ID 列表，由 Trigger Authoring 发布。"
    if name.endswith("VfxId"):
        return f"{translate(name[:-5])}特效资源 ID，引用特效资源表。"
    if name.endswith("Ms"):
        return f"{translate(name[:-2])}，单位毫秒。"
    if name.endswith("Deg"):
        return f"{translate(name[:-3])}，单位度。"
    if name.endswith("Id") or name.endswith("Ids"):
        return f"{translate(name)}；填写对应配置或编辑器产物的 ID。"
    label = translate(name)
    if kind == "bool":
        return f"{label}：true 为启用，false 为关闭。"
    if kind.startswith("(list") or kind.startswith("list,"):
        return f"{label}列表；按列类型填写，顺序可能影响运行时行为。"
    if kind not in ("int", "float", "string", "bool"):
        return f"{label}；按 {kind} 枚举或 bean 定义填写。"
    return f"{label}；用于{sheet}配置。"


def translate(name):
    words = re.findall(r"[A-Z]+(?=[A-Z][a-z]|$)|[A-Z]?[a-z]+|[0-9]+", name)
    if not words:
        raise ValueError(f"cannot describe field {name}")
    missing = [word for word in words if word not in TOKENS]
    if missing:
        raise ValueError(f"field {name} has untranslated tokens: {missing}")
    return "".join(TOKENS[word] for word in words)


def annotate_sheet(sheet, apply):
    nested = sheet.cell(3, 1).value == "##var"
    note_row = 4 if nested else 3
    if nested and sheet.cell(4, 1).value != "##" and apply:
        sheet.insert_rows(4)
    if not nested and sheet.cell(3, 1).value != "##":
        raise ValueError(f"{sheet.title}: unexpected header row")
    if nested and sheet.cell(4, 1).value != "##" and not apply:
        raise ValueError(f"{sheet.title}: missing nested ## note row")
    if apply:
        sheet.cell(note_row, 1, "##")
        sheet.row_dimensions[note_row].height = 58
    previous_root = ""
    count = 0
    for col in range(2, sheet.max_column + 1):
        header = sheet.cell(1, col).value
        if header:
            previous_root = str(header).lstrip("*")
        member = sheet.cell(3, col).value if nested else None
        field = str(member or header or "")
        if not field:
            continue
        kind = str(sheet.cell(2, col).value or "")
        if member:
            kind = "string"  # The bean's member type is defined in strong_types.xml.
        description = describe(field.lstrip("*"), kind, sheet.title)
        if member:
            description = f"{previous_root}.{field}：{description}"
        if apply:
            cell = sheet.cell(note_row, col)
            cell.value = description
            cell.font = copy(sheet.cell(1, col).font)
            cell.fill = NOTE_FILL
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        else:
            current = sheet.cell(note_row, col).value
            if not current:
                raise ValueError(f"{sheet.title}!{sheet.cell(note_row, col).coordinate}: missing note for {field}")
        count += 1
    return count


def annotate_control_sheet(sheet, apply):
    if sheet.cell(3, 1).value != "##":
        raise ValueError(f"{sheet.title}: unexpected control-sheet note row")
    count = 0
    for col in range(2, sheet.max_column + 1):
        header = sheet.cell(1, col).value
        note = sheet.cell(3, col)
        if not header and not note.value:
            continue
        if not note.value:
            if not header or header not in SPECIAL:
                raise ValueError(f"{sheet.title}!{note.coordinate}: missing control-field description")
            if apply:
                note.value = SPECIAL[header]
                note.fill = NOTE_FILL
                note.alignment = Alignment(vertical="center", wrap_text=True)
        count += 1
    return count


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    sheets = fields = 0
    for path in sorted(DATAS.glob("*.xlsx")):
        book = load_workbook(path)
        for sheet in book:
            if path.name.startswith("__"):
                fields += annotate_control_sheet(sheet, args.apply)
            else:
                fields += annotate_sheet(sheet, args.apply)
            sheets += 1
        if args.apply:
            book.save(path)
    print(f"validated {fields} field notes across {sheets} production sheets")


if __name__ == "__main__":
    main()
