"""Preview modifier template dependents in the authoritative effect workbook."""

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from modifier_parameters import domain_name_for_parameter_id  # noqa: E402

DEFAULT_WORKBOOK = Path(__file__).resolve().parents[1] / "Production/Datas/J-效果表.xlsx"


def read_workbook(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    result = {}
    markers = {
        # Mount rows only carry DefinitionId; the semantics live in modifier_definitions.
        "modifier_templates": {"DefinitionId", "Selectors"},
        "modifier_definitions": {"Target", "Value", "Key"},
        # The parameter catalog is authored with Luban field names plus author-facing aliases.
        "modifier_parameters": {"Key", "Category", "SubCategory"},
        "buffs": {"ModifierIds", "MaxStacks"},
        "continuous_processes": {"ModifierIds", "IntervalTriggerIds"},
    }
    for sheet in workbook:
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        columns = {name: index for index, name in enumerate(header) if name}
        matches = [name for name, fields in markers.items() if fields <= columns.keys()]
        if not matches:
            continue
        if len(matches) != 1:
            raise ValueError(f"Ambiguous modifier sheet {sheet.title}")
        name = matches[0]
        next(rows)  # Luban types
        next(rows)  # descriptions
        records = []
        for values in rows:
            identity = values[columns["Id"]]
            if type(identity) is not int:
                continue
            record = {field: values[index] for field, index in columns.items()}
            if name in ("buffs", "continuous_processes"):
                raw = record.get("ModifierIds") or ""
                record["ModifierIds"] = [int(part) for part in str(raw).split(",") if part.strip()]
            records.append(record)
        result[name] = records
    workbook.close()
    missing = markers.keys() - result.keys()
    if missing:
        raise ValueError(f"Missing modifier sheets: {sorted(missing)}")
    return result


def build_report(templates, buffs, processes, parameters=(), definitions=()):
    by_id = {row["Id"]: row for row in templates}
    if len(by_id) != len(templates):
        raise ValueError("Duplicate modifier template ID")
    dependents = {identity: [] for identity in by_id}
    # The catalog no longer stores Domain; it is implied by each parameter's ID segment.
    parameter_domains = {row["Id"]: domain_name_for_parameter_id(row["Id"])
                         for row in parameters}
    by_definition = {row["Id"]: row for row in definitions}
    for table, rows in (("buffs", buffs), ("continuous_processes", processes)):
        for row in rows:
            for identity in row.get("ModifierIds") or []:
                if identity not in dependents:
                    raise ValueError(f"{table}/{row['Id']}: missing modifier template {identity}")
                dependents[identity].append({"table": table, "id": row["Id"]})

    entries = []
    coverage = {}
    for identity, row in sorted(by_id.items()):
        references = sorted(dependents[identity], key=lambda item: (item["table"], item["id"]))
        definition = by_definition.get(row.get("DefinitionId"), {})
        # Definitions are authored as the `Kind|Id` bean text rather than a nested object.
        target = str(definition.get("Target") or "0|0").split("|")
        target_kind, target_id = int(target[0]), int(target[1]) if len(target) > 1 else 0
        domain = ("Attribute" if target_kind in (1, "Attribute") else
                  "Skill" if target_kind in (5, "SkillSlotCandidate") else
                  parameter_domains.get(target_id, "Unknown"))
        coverage[str(domain)] = coverage.get(str(domain), 0) + 1
        entries.append({
            "id": identity,
            "domain": domain,
            "definitionId": row.get("DefinitionId"),
            "targetKind": target_kind,
            "targetId": target_id,
            "usedBy": references,
            "shared": len(references) > 1,
            "unused": not references,
        })
    return {
        "schemaVersion": 1,
        "templateCount": len(entries),
        "domainCoverage": dict(sorted(coverage.items())),
        "sharedIds": [entry["id"] for entry in entries if entry["shared"]],
        "unusedIds": [entry["id"] for entry in entries if entry["unused"]],
        "templates": entries,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    tables = read_workbook(args.workbook)
    report = build_report(tables["modifier_templates"], tables["buffs"],
                          tables["continuous_processes"], tables["modifier_parameters"],
                          tables["modifier_definitions"])
    output = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(output, encoding="utf-8")
    else:
        print(output, end="")


if __name__ == "__main__":
    main()
