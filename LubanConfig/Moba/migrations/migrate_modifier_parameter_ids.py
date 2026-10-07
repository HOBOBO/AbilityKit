"""One-time domain allocation for the published modifier parameter catalog."""

from pathlib import Path
import sys

from openpyxl import load_workbook
from openpyxl.comments import Comment

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from modifier_parameters import EMITTER_LEGACY_IDS, PARAMETER_ID_MIGRATION


WORKBOOK = Path(__file__).resolve().parents[1] / "Production/Datas/J-效果表.xlsx"
EMITTER_DEFINITION_IDS = {500001: 400001, 500002: 400002}


def column(sheet, name):
    return next(cell.column for cell in sheet[1] if cell.value == name)


def migrate():
    workbook = load_workbook(WORKBOOK)
    parameters = workbook["修饰参数"]
    id_column = column(parameters, "Id")
    domain_column = column(parameters, "Domain")
    current = {parameters.cell(row, id_column).value for row in range(4, parameters.max_row + 1)
               if isinstance(parameters.cell(row, id_column).value, int)}
    if current not in (set(PARAMETER_ID_MIGRATION), set(PARAMETER_ID_MIGRATION.values())):
        raise ValueError("modifier parameter IDs differ from the audited migration map")
    definitions = workbook["修饰器定义"]
    target_kind_column = column(definitions, "TargetKind")
    target_id_column = column(definitions, "TargetId")
    if current == set(PARAMETER_ID_MIGRATION):
        for row in range(4, parameters.max_row + 1):
            old_id = parameters.cell(row, id_column).value
            if not isinstance(old_id, int):
                continue
            parameters.cell(row, id_column).value = PARAMETER_ID_MIGRATION[old_id]
            if old_id in EMITTER_LEGACY_IDS:
                parameters.cell(row, domain_column).value = "Emitter"
        values = [tuple(cell.value for cell in parameters[row]) for row in range(4, parameters.max_row + 1)]
        for row, record in enumerate(sorted(values, key=lambda item: item[id_column - 1]), 4):
            for col, value in enumerate(record, 1):
                parameters.cell(row, col).value = value
        note = parameters.cell(3, id_column)
        note.value = "稳定参数 ID；现有行已按 Domain 分段，旧 1-16 停用。详见批注及 Docs/MobaModifierAuthoring.md。"
        note.comment = Comment(
            "技能 10000-19999；Buff 20000-29999；发射器 30000-39999；子弹 40000-49999；"
            "AOE 50000-59999；召唤 60000-69999；伤害 70000-79999。"
            "属性目标引用独立 attr_types.Id。旧参数 ID 1-16 迁移后不再分配。",
            "AbilityKit")
        for row in range(4, definitions.max_row + 1):
            if definitions.cell(row, target_kind_column).value not in ("SkillParameter", "ParameterReferenceCandidate"):
                continue
            target = definitions.cell(row, target_id_column)
            target.value = PARAMETER_ID_MIGRATION[target.value]

    definition_id_column = column(definitions, "Id")
    current_definitions = {definitions.cell(row, definition_id_column).value
                           for row in range(4, definitions.max_row + 1)}
    if set(EMITTER_DEFINITION_IDS) <= current_definitions:
        for row in range(4, definitions.max_row + 1):
            identity = definitions.cell(row, definition_id_column)
            if identity.value in EMITTER_DEFINITION_IDS:
                identity.value = EMITTER_DEFINITION_IDS[identity.value]
        values = [tuple(cell.value for cell in definitions[row]) for row in range(4, definitions.max_row + 1)]
        for row, record in enumerate(sorted(values, key=lambda item: item[definition_id_column - 1]), 4):
            for col, value in enumerate(record, 1):
                definitions.cell(row, col).value = value
        buffs = workbook["Buff"]
        reference_column = column(buffs, "ModifierIds")
        for row in range(4, buffs.max_row + 1):
            cell = buffs.cell(row, reference_column)
            if cell.value:
                ids = [int(part.strip()) for part in str(cell.value).split(",")]
                cell.value = ",".join(str(EMITTER_DEFINITION_IDS.get(identity, identity)) for identity in ids)
    elif not set(EMITTER_DEFINITION_IDS.values()) <= current_definitions:
        raise ValueError("emitter modifier definition IDs differ from the audited migration map")
    workbook.save(WORKBOOK)


if __name__ == "__main__":
    migrate()
