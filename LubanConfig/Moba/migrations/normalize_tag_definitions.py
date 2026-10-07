"""One-time migration of the production tag dictionary and legacy rule aliases."""

from copy import copy
from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path(__file__).resolve().parents[1] / "Production/Datas/J-标签表.xlsx"
HEROES = ("LianPo", "XiaoQiao", "ZhaoYun", "MoZi", "DaJi", "YingZheng")
HERO_NAMES = ("廉颇", "小乔", "赵云", "墨子", "妲己", "嬴政")
STATE_NAMES = {
    1001: "控制免疫", 1002: "霸体", 1006: "减速", 1013: "睡眠",
    1014: "魅惑", 1015: "失能", 1016: "恐惧", 1017: "伤害免疫",
    1018: "禁锢", 1019: "沉默", 1020: "眩晕", 1021: "压制",
    1022: "不可选取",
}
RULE_NAMES = {
    2001: "不可选取", 2002: "伤害免疫", 2003: "禁止施法", 2004: "硬控",
    2005: "禁止移动", 2006: "控制免疫", 2007: "不可受控制",
    2008: "禁止攻击", 10000001: "控制位移霸体阻断",
}


def normalize_rule_names(book):
    sheet = book["标签规则组"]
    for row in list(sheet.rows)[3:]:
        if row[1].value in RULE_NAMES:
            row[2].value = RULE_NAMES[row[1].value]


def normalize_rule_node_ids(book):
    sheet = book["标签规则节点"]
    data = [row for row in list(sheet.rows)[3:] if isinstance(row[1].value, int)]
    if not data or all(row[1].value // 100 == row[2].value for row in data):
        return
    counters = {}
    new_ids = {}
    for row in data:
        group_id = row[2].value
        counters[group_id] = counters.get(group_id, 0) + 1
        if counters[group_id] > 99:
            raise ValueError(f"Too many tag rule nodes in group {group_id}")
        new_ids[row[1].value] = group_id * 100 + counters[group_id]
    for row in data:
        old_id = row[1].value
        row[1].value = new_ids[old_id]
        if row[3].value:
            row[3].value = new_ids[row[3].value]
    for row in list(book["标签规则组"].rows)[3:]:
        row[3].value = new_ids[row[3].value]


def replace_rows(sheet, rows):
    styles = [copy(sheet.cell(4, col)._style) for col in range(1, sheet.max_column + 1)]
    if sheet.title == "标签定义":
        styles[4] = copy(styles[2])
        styles[5] = copy(styles[1])
    sheet.delete_rows(4, max(0, sheet.max_row - 3))
    for number, values in enumerate(rows, 4):
        for col, value in enumerate(values, 1):
            cell = sheet.cell(number, col, value)
            cell._style = styles[col - 1]


def main():
    book = load_workbook(WORKBOOK)
    definitions = book["标签定义"]
    if definitions.cell(1, 5).value == "DisplayName":
        normalize_rule_names(book)
        normalize_rule_node_ids(book)
        definitions.cell(3, 2, "稳定 ID：90 是可选技能标签根；1000-1099 是现有战斗标签冻结段；1100-1499 为状态，1500-1999 为负面效果；2000-2999 为按需技能标签。")
        definitions.auto_filter.ref = f"B1:F{definitions.max_row}"
        definitions.column_dimensions["C"].width = 38
        definitions.column_dimensions["E"].width = 28
        definitions.column_dimensions["F"].width = 18
        nodes = book["标签规则节点"]
        nodes.auto_filter.ref = f"B1:I{nodes.max_row}"
        book.save(WORKBOOK)
        print("updated migrated tag workbook formatting")
        return
    original = [list(row) for row in list(definitions.values)[3:]]
    if definitions.cell(1, 5).value is not None:
        raise RuntimeError("Tag dictionary has an unexpected fifth column")

    definitions.cell(1, 5, "DisplayName")
    definitions.cell(2, 5, "string")
    definitions.cell(3, 5, "策划可读中文名称；运行时比较使用 Name。")
    definitions.cell(1, 6, "ParentId")
    definitions.cell(2, 6, "int")
    definitions.cell(3, 6, "直接父标签 ID；根节点填 0，Name 必须为父 Name 加点号和自身段。")
    for row in (1, 2, 3):
        for col in (5, 6):
            definitions.cell(row, col)._style = copy(definitions.cell(row, 3 if col == 5 else 2)._style)
    definitions.cell(3, 2, "稳定 ID：90 是可选技能标签根；1000-1099 是现有战斗标签冻结段；1100-1499 为状态，1500-1999 为负面效果；2000-2999 为按需技能标签。")
    definitions.cell(3, 3, "运行时唯一 ASCII 点分标签路径；中文名称填写 DisplayName。")

    rows = []
    for _, identity, name, namespace in original:
        if namespace == "LegacyNumeric":
            rows.append((None, identity, name, namespace, f"旧数字标签 {identity}", 0))
        elif namespace == "GameplayState" and (name == "Debuff.Slow" or name.startswith("State.")):
            rows.append((None, identity, name, namespace, STATE_NAMES[identity],
                         1040 if name.startswith("Debuff.") else 1000))
        elif namespace == "SkillSelection":
            hero_index, slot = divmod(identity, 100)
            if not 1 <= hero_index <= len(HEROES) or not 1 <= slot <= 3:
                raise ValueError(f"Unexpected skill tag ID: {identity}")
            rows.append((None, identity, f"Skill.Hero.{HEROES[hero_index - 1]}.S{slot}",
                         namespace, name, hero_index * 100))

    rows.extend(((None, 90, "Skill", "SkillSelection", "技能分类", 0),
                 (None, 91, "Skill.Hero", "SkillSelection", "英雄技能", 90),
                 (None, 1000, "State", "GameplayState", "战斗状态", 0),
                 (None, 1040, "Debuff", "GameplayState", "负面效果", 0)))
    for index, (hero, display) in enumerate(zip(HEROES, HERO_NAMES), 1):
        rows.append((None, index * 100, f"Skill.Hero.{hero}", "SkillSelection", display, 91))
    replace_rows(definitions, sorted(rows, key=lambda row: row[1]))

    kept_ids = {row[1] for row in rows}
    nodes = book["标签规则节点"]
    original_nodes = [list(row) for row in list(nodes.values)[3:]]
    filtered = [row for row in original_nodes if row[6] == 0 or row[6] in kept_ids]
    replace_rows(nodes, filtered)
    normalize_rule_names(book)
    normalize_rule_node_ids(book)

    definitions.auto_filter.ref = f"B1:F{definitions.max_row}"
    definitions.column_dimensions["C"].width = 38
    definitions.column_dimensions["E"].width = 28
    definitions.column_dimensions["F"].width = 18
    nodes.auto_filter.ref = f"B1:I{nodes.max_row}"
    book.save(WORKBOOK)
    print(f"tag definitions: {len(original)} -> {len(rows)}; rule nodes: {len(original_nodes)} -> {len(filtered)}")


if __name__ == "__main__":
    main()
