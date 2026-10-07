"""One-time migration of the MOBA modifier parameter catalog to Excel."""

from copy import copy
from pathlib import Path

from openpyxl import load_workbook


DATAS = Path(__file__).resolve().parents[1] / "Production/Datas"
ROWS = (
    (1, "Projectile", "LauncherId", 30, 1, "ReferenceId", 4, False, "Action"),
    (2, "Projectile", "ProjectileId", 30, 2, "ReferenceId", 4, False, "Action"),
    (3, "Projectile", "CountPerShot", 30, 3, "Integer", 15, True, "Launch"),
    (4, "Projectile", "FanAngleDeg", 30, 4, "Float", 15, True, "Launch"),
    (5, "Projectile", "DurationMs", 30, 5, "Integer", 15, True, "Launch"),
    (6, "Skill", "SkillId", 40, 1, "ReferenceId", 4, False, "Action"),
    (7, "Summon", "SummonId", 40, 2, "ReferenceId", 4, False, "Action"),
    (8, "Skill", "ResourceCost", 40, 3, "Integer", 15, True, "Action"),
    (9, "Skill", "CooldownMs", 40, 4, "Integer", 15, True, "Action"),
    (10, "Skill", "CastRange", 40, 5, "Float", 15, True, "Action"),
    (11, "Buff", "DurationMs", 32, 1, "Integer", 15, True, "BuffApply"),
    (12, "Projectile", "Speed", 30, 6, "Float", 15, True, "Launch"),
    (13, "Aoe", "Radius", 31, 1, "Float", 15, True, "AreaSpawn"),
    (14, "Aoe", "DurationMs", 31, 2, "Integer", 15, True, "AreaSpawn"),
)


def main():
    path = DATAS / "J-效果表.xlsx"
    workbook = load_workbook(path)
    if "修饰参数" in workbook:
        del workbook["修饰参数"]
    source = workbook["修饰器定义"]
    sheet = workbook.create_sheet("修饰参数", workbook.sheetnames.index("修饰器定义"))
    for row_number, marker in ((1, "##var"), (2, "##type"), (3, "##")):
        sheet.cell(row_number, 1, marker)._style = copy(source.cell(row_number, 1)._style)
    for index, (name, kind, comment) in enumerate((
        ("Id", "int", "稳定参数 ID；发布后不可重用或改变含义。"),
        ("Domain", "ModifierParameterDomain", "参数所属领域，决定生成代码中的分组。"),
        ("Key", "string", "领域内唯一的 C# 参数名；发布后不可改名。"),
        ("Category", "int", "ModifierKey 类别位，范围 1-255；发布后不可改动。"),
        ("SubCategory", "int", "ModifierKey 子类别位，范围 1-255；发布后不可改动。"),
        ("ValueKind", "ModifierParameterValueKind", "ReferenceId 只允许 Override；Integer/Float 为数值。"),
        ("OpMask", "int", "允许的操作位：Add=1、Mul=2、Override=4、PercentAdd=8。"),
        ("TriggerWritable", "bool", "触发器的 add_skill_param_modifier 动作是否允许写入。"),
        ("ResolveAt", "ModifierParameterResolveAt", "参数取值时机；已有实例不会追溯修改。"),
    ), 1):
        for row_number, value in ((1, name), (2, kind), (3, comment)):
            cell = sheet.cell(row_number, index + 1, value)
            cell._style = copy(source.cell(row_number, min(index + 1, source.max_column))._style)
        sheet.column_dimensions[sheet.cell(1, index + 1).column_letter].width = 25 if index != 3 else 22
    for row_number, values in enumerate(ROWS, 4):
        for index, value in enumerate(values, 1):
            cell = sheet.cell(row_number, index + 1, value)
            cell._style = copy(source.cell(4, min(index + 1, source.max_column))._style)
    workbook.save(path)

    path = DATAS / "__tables__.xlsx"
    catalog = load_workbook(path)
    sheet = catalog.active
    if not any(row[1] == "ModifierParameters" for row in sheet.values):
        sheet.append((None, "ModifierParameters", "DRModifierParameters", True,
                      "修饰参数@J-效果表.xlsx", "Id", None, "c,s,e", None, None,
                      "modifier_parameters"))
    catalog.save(path)


if __name__ == "__main__":
    main()
