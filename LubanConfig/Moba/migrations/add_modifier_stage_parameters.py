"""Register the first stage-specific emitter, projectile and area parameters."""

from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path(__file__).resolve().parents[1] / "Production/Datas/J-效果表.xlsx"
PARAMETERS = (
    (30005, "Emitter", "IntervalMs", 30, 7, "Integer", 15, False, "Launch"),
    (40003, "Projectile", "LifetimeMs", 30, 8, "Integer", 15, False, "ProjectileSpawn"),
    (50003, "Aoe", "IntervalMs", 31, 3, "Integer", 15, False, "AreaSpawn"),
)


def main():
    workbook = load_workbook(WORKBOOK)
    sheets = [sheet for sheet in workbook
              if any(cell.value == "ResolveAt" for cell in sheet[1])]
    if len(sheets) != 1:
        raise ValueError("Expected exactly one modifier parameter sheet")
    sheet = sheets[0]
    existing = {sheet.cell(row, 2).value: tuple(sheet.cell(row, col).value
                                                  for col in range(2, 11))
                for row in range(4, sheet.max_row + 1)
                if isinstance(sheet.cell(row, 2).value, int)}
    for parameter in PARAMETERS:
        if parameter[0] in existing and existing[parameter[0]] != parameter:
            raise ValueError(f"Modifier parameter {parameter[0]} already has different data")
        existing[parameter[0]] = parameter
    for row, parameter in enumerate(sorted(existing.values()), 4):
        for col, value in enumerate(parameter, 2):
            sheet.cell(row, col, value)
    workbook.save(WORKBOOK)


if __name__ == "__main__":
    main()
