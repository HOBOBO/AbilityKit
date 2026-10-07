"""One-time, idempotent migration of modifier selectors in the production workbook."""

from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path(__file__).resolve().parents[1] / "Production/Datas/J-效果表.xlsx"
SELECTORS = (
    "SelectorSlot",
    "SelectorBaseSkillId",
    "SelectorEffectiveSkillId",
    "SelectorLauncherTemplateId",
    "SelectorProjectileTemplateId",
    "SelectorAreaTemplateId",
    "SelectorSummonTemplateId",
    "SelectorBuffTemplateId",
    "CandidateSkillId",
    "CandidateReferenceId",
    "FormulaConstant",
    "FormulaLevelCoefficient",
    "FormulaTerms",
)


def sheet_with_column(workbook, column):
    matches = [sheet for sheet in workbook if any(cell.value == column for cell in sheet[1])]
    if len(matches) != 1:
        raise ValueError(f"Expected one sheet with {column}, found {len(matches)}")
    return matches[0]


def main():
    workbook = load_workbook(WORKBOOK)
    templates = sheet_with_column(workbook, "MagnitudeSourceType")
    existing = {cell.value for cell in templates[1]}
    for field in SELECTORS:
        if field in existing:
            continue
        column = templates.max_column + 1
        templates.cell(1, column, field)
        field_type = "(list#sep=,),ModifierFormulaTerm" if field == "FormulaTerms" else \
            "float" if field in ("FormulaConstant", "FormulaLevelCoefficient") else "int"
        templates.cell(2, column, field_type)
        templates.cell(3, column, "Bounded snapshot formula" if field.startswith("Formula") else
                       "0 means any; positive ID matches the stage record")
        for row in range(4, templates.max_row + 1):
            if isinstance(templates.cell(row, 2).value, int):
                templates.cell(row, column, "" if field == "FormulaTerms" else 0)

    parameters = sheet_with_column(workbook, "ResolveAt")
    keys = {parameters.cell(row, 4).value: row for row in range(4, parameters.max_row + 1)}
    for key, stage in (("CountPerShot", "Emit"), ("FanAngleDeg", "Emit"),
                       ("Speed", "ProjectileSpawn")):
        parameters.cell(keys[key], 10, stage)
    workbook.save(WORKBOOK)


if __name__ == "__main__":
    main()
