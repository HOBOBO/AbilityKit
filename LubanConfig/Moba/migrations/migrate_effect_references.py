"""One-time extraction of MOBA modifier and tag definitions into ID tables."""

import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

from migrate_strong_types import ENUMS


ROOT = Path(__file__).resolve().parents[3]
DATA = ROOT / "LubanConfig/Moba/Production/Datas"
BASELINE = ROOT / "Unity/Packages/com.abilitykit.demo.moba.view.runtime/Resources/moba"
WORKBOOK = DATA / "J-效果表.xlsx"
MODIFIER_FIELDS = [
    ("TargetKind", "MobaModifierTargetKind"), ("TargetId", "int"),
    ("AttrTypeId", "int"), ("Op", "ModifierOp"), ("Value", "float"),
    ("Priority", "int"), ("MagnitudeSourceType", "MagnitudeSourceType"),
    ("MagnitudeBaseValue", "float"), ("MagnitudeCoefficient", "float"),
    ("MagnitudeDuration", "float"), ("MagnitudeDecayType", "int"),
    ("MagnitudeAttributeTypeId", "int"), ("MagnitudeContextKey", "string"),
    ("MagnitudeCurve", "(list#sep=,),float"),
    ("EvaluationPolicy", "ModifierEvaluationPolicy"),
]
TAG_FIELDS = ("RequiredTags", "BlockedTags", "GrantTags", "RemoveTags")
CONTINUOUS_TAG_FIELDS = (
    "ActivationRequiredTags", "ActivationBlockedTags", "ApplicationTags",
    "RemovalRequiredTags", "RemovalBlockedTags", "OngoingRequiredTags",
    "OngoingBlockedTags", "RemovalTags",
)


def baseline(name):
    return json.loads((BASELINE / f"{name}.json").read_text(encoding="utf-8-sig"))


def enum_name(value, name):
    if value is None:
        return None
    values = ENUMS[name]
    return values[value] if isinstance(values, list) else values[value]


def replace_sheet(book, title, changes, values_by_id):
    source = book[title]
    position = book.sheetnames.index(title)
    target = book.create_sheet(title + "_new", position)
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    fields = [(c.value, source.cell(2, c.column).value, c.column)
              for c in source[1] if c.value and c.value != "##var"]
    output = []
    for field, kind, column in fields:
        if field == "*Modifiers":
            output.append(("ModifierIds", "(list#sep=,),int", None))
        elif field in changes:
            output.append((changes[field], "(list#sep=,),int", None))
        else:
            output.append((field, kind, column))
    target.cell(1, 1, "##var")
    target.cell(2, 1, "##type")
    target.cell(3, 1, "##")
    target.cell(3, 2, "ID references; edit definition sheets below.")
    for index, (field, kind, old_column) in enumerate(output, 2):
        for row, value in ((1, field), (2, kind)):
            cell = target.cell(row, index, value)
            if old_column:
                original = source.cell(row, old_column)
                cell._style = copy(original._style)
        if old_column:
            target.column_dimensions[target.cell(1, index).column_letter].width = \
                source.column_dimensions[source.cell(1, old_column).column_letter].width
    source_headers = {name: column for name, _, column in fields}
    destination_row = 4
    for source_row in range(4, source.max_row + 1):
        identity = source.cell(source_row, source_headers["Id"]).value
        if identity is None:
            continue
        mapped = values_by_id[identity]
        for index, (field, _, old_column) in enumerate(output, 2):
            value = mapped.get(field) if old_column is None else source.cell(source_row, old_column).value
            cell = target.cell(destination_row, index, value)
            if old_column:
                cell._style = copy(source.cell(source_row, old_column)._style)
        destination_row += 1
    book.remove(source)
    target.title = title


def add_definition_sheet(book, title, columns, rows):
    sheet = book.create_sheet(title)
    sheet.append(["##var", *[name for name, _ in columns]])
    sheet.append(["##type", *[kind for _, kind in columns]])
    sheet.append(["##", "Reusable definition; references are validated during export."])
    for row in rows:
        sheet.append([None, *row])
    sheet.column_dimensions["B"].width = 18
    for column in range(3, len(columns) + 2):
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = 22


def register(book, name, value_type, sheet, output):
    table = book.active
    table.append([None, name, value_type, True,
                  f"{sheet}@J-效果表.xlsx", "Id", None, "c,s,e", None, None, output])


def migrate():
    book = load_workbook(WORKBOOK)
    if "修饰器定义" in book or "标签定义" in book:
        raise ValueError("Effect references have already been migrated")
    modifiers = []
    references = {}
    next_id = 1
    for table_name, kind in (("buffs", "Buff"), ("continuous_processes", "ContinuousProcess")):
        for record in baseline(table_name):
            ids = []
            for modifier in record.get("Modifiers") or []:
                ids.append(next_id)
                values = []
                for field, value_type in MODIFIER_FIELDS:
                    value = modifier.get(field)
                    if value_type in ENUMS:
                        value = enum_name(value, value_type)
                    elif isinstance(value, list):
                        value = ",".join(str(item) for item in value)
                    values.append(value)
                modifiers.append((next_id, kind, *values))
                next_id += 1
            references[table_name, record["Id"]] = {"ModifierIds": ",".join(map(str, ids)) or None}

    tag_names = {1: "1", 2: "2", 3: "3", 1001: "State.ControlImmune", 1002: "State.SuperArmor"}
    tag_ids = {name: identity for identity, name in tag_names.items()}
    for table_name, source_fields in (("buffs", ("TagNames", "DispelBlockedByTagNames")),
                                      ("continuous_processes", ("Tags",)),
                                      ("tag_templates", TAG_FIELDS),
                                      ("continuous_tag_templates", CONTINUOUS_TAG_FIELDS)):
        for record in baseline(table_name):
            mapped = references.setdefault((table_name, record["Id"]), {})
            for field in source_fields:
                names = record.get(field) or []
                mapped[field.replace("TagNames", "TagIds").replace("Tags", "TagIds")] = \
                    ",".join(str(tag_ids[str(name)]) for name in names) or None

    replace_sheet(book, "Buff", {"TagNames": "TagIds",
                                  "DispelBlockedByTagNames": "DispelBlockedByTagIds"},
                  {key[1]: value for key, value in references.items() if key[0] == "buffs"})
    replace_sheet(book, "持续流程", {"Tags": "TagIds"},
                  {key[1]: value for key, value in references.items() if key[0] == "continuous_processes"})
    replace_sheet(book, "标签模板", {field: field[:-4] + "TagIds" for field in TAG_FIELDS},
                  {key[1]: value for key, value in references.items() if key[0] == "tag_templates"})
    replace_sheet(book, "持续标签模板", {field: field[:-4] + "TagIds" for field in CONTINUOUS_TAG_FIELDS},
                  {key[1]: value for key, value in references.items() if key[0] == "continuous_tag_templates"})
    add_definition_sheet(book, "修饰器定义", [("Id", "int"), ("Kind", "ModifierTemplateKind"),
                                         *MODIFIER_FIELDS], modifiers)
    add_definition_sheet(book, "标签定义", [("Id", "int"), ("Name", "string")],
                         [(identity, name) for identity, name in tag_names.items()])
    book.save(WORKBOOK)

    registry_path = DATA / "__tables__.xlsx"
    registry = load_workbook(registry_path)
    register(registry, "ModifierTemplates", "DRModifierTemplates", "修饰器定义", "modifier_templates")
    register(registry, "TagDefinitions", "DRTagDefinitions", "标签定义", "tag_definitions")
    registry.save(registry_path)


if __name__ == "__main__":
    migrate()
