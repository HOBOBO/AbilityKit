"""One-time migration of modifier definition IDs in the production effect workbook."""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment


WORKBOOK = Path(__file__).resolve().parents[1] / "Production/Datas/J-效果表.xlsx"
ID_MAP = {
    1: 400001, 2: 400002,
    3: 100001, 4: 100002, 5: 100003, 6: 100004, 7: 100005,
    8: 100006, 9: 100007, 10: 100008, 11: 100009, 12: 100010,
    13: 100011, 14: 100012, 15: 100013, 16: 100014,
    18: 100015, 19: 100016, 20: 300001, 21: 100017,
    23: 100018, 24: 100019, 26: 100020, 27: 100021, 28: 100022,
}


def column(sheet, name):
    return next(cell.column for cell in sheet[1] if cell.value == name)


def migrate():
    workbook = load_workbook(WORKBOOK)
    definitions = workbook["修饰器定义"]
    id_column = column(definitions, "Id")
    current = {definitions.cell(row, id_column).value for row in range(4, definitions.max_row + 1)
               if isinstance(definitions.cell(row, id_column).value, int)}
    if current == set(ID_MAP.values()):
        return
    if current != set(ID_MAP):
        raise ValueError("modifier definition IDs differ from the audited migration map")
    for row in range(4, definitions.max_row + 1):
        cell = definitions.cell(row, id_column)
        if isinstance(cell.value, int):
            cell.value = ID_MAP[cell.value]
    values = [tuple(cell.value for cell in definitions[row]) for row in range(4, definitions.max_row + 1)]
    for row, record in enumerate(sorted(values, key=lambda item: item[id_column - 1]), 4):
        for col, value in enumerate(record, 1):
            definitions.cell(row, col).value = value
    definitions.cell(3, id_column).value = "稳定修饰器定义 ID；按目标类型分段，详见批注及 Docs/MobaModifierAuthoring.md。"
    definitions.cell(3, id_column).comment = Comment(
        "属性 100001-199999；Buff 200001-299999；技能 300001-399999；发射器 400001-499999；"
        "子弹 500001-599999；AOE 600001-699999；召唤 700001-799999；伤害 800001-899999。"
        "按 TargetKind 和修饰参数 Domain 归类，不按提供者 Kind 归类；旧 ID 1-28 已迁移，不再分配。",
        "AbilityKit")
    for sheet_name in ("Buff", "持续流程"):
        sheet = workbook[sheet_name]
        reference_column = column(sheet, "ModifierIds")
        for row in range(4, sheet.max_row + 1):
            cell = sheet.cell(row, reference_column)
            if not cell.value:
                continue
            ids = [int(part.strip()) for part in str(cell.value).split(",")]
            cell.value = ",".join(str(ID_MAP[identity]) for identity in ids)
    workbook.save(WORKBOOK)


if __name__ == "__main__":
    migrate()
