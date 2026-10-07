"""Register stage-consumed projectile hit and summon inheritance parameters."""

from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path(__file__).resolve().parents[1] / "Production/Datas/J-效果表.xlsx"
PARAMETERS = (
    (40004, "Projectile", "PierceHits", 30, 9, "Integer", 15, False, "ProjectileSpawn"),
    (60003, "Summon", "InheritRatioScale", 40, 7, "Float", 15, False, "SummonCreate"),
)


def main():
    workbook = load_workbook(WORKBOOK)
    sheets = [sheet for sheet in workbook
              if any(cell.value == "ResolveAt" for cell in sheet[1])]
    if len(sheets) != 1:
        raise ValueError("Expected exactly one modifier parameter sheet")
    sheet = sheets[0]
    existing = {sheet.cell(row, 2).value: tuple(sheet.cell(row, col).value
                                                  for col in range(2, 11))
                for row in range(4, sheet.max_row + 1)
                if isinstance(sheet.cell(row, 2).value, int)}
    for parameter in PARAMETERS:
        if parameter[0] in existing and existing[parameter[0]] != parameter:
            raise ValueError(f"Modifier parameter {parameter[0]} already differs")
        existing[parameter[0]] = parameter
    for row, parameter in enumerate(sorted(existing.values()), 4):
        for col, value in enumerate(parameter, 2):
            sheet.cell(row, col, value)
    workbook.save(WORKBOOK)


if __name__ == "__main__":
    main()
