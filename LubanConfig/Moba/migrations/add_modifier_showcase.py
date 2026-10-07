"""Install isolated modifier examples in the authoritative MOBA workbooks."""

import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


DATA = Path(__file__).resolve().parents[1] / "Production/Datas"
EFFECTS = DATA / "J-效果表.xlsx"
SUMMONS = DATA / "J-召唤物表.xlsx"

TEMPLATES = (
    (100023, {"TargetKind": "Attribute", "TargetId": 17, "Op": "PercentAdd",
              "Value": 0.4, "MagnitudeSourceType": "TimeDecay",
              "MagnitudeDuration": 3, "MagnitudeCurve": "1,0.5,0.1"}),
    (100024, {"TargetKind": "Attribute", "TargetId": 4, "Op": "Add",
              "Value": 0, "MagnitudeSourceType": "LinearFormula",
              "FormulaConstant": 5, "FormulaLevelCoefficient": 2,
              "EvaluationPolicy": "OnApplySnapshot"}),
    (200001, {"TargetKind": "SkillParameter", "TargetId": 20001,
              "Op": "Add", "Value": 1000, "SelectorBuffTemplateId": 10020001}),
    (500001, {"TargetKind": "SkillParameter", "TargetId": 40002,
              "Op": "Mul", "Value": 1.25, "SelectorBaseSkillId": 10020101,
              "SelectorProjectileTemplateId": 30020101}),
    (500002, {"TargetKind": "SkillParameter", "TargetId": 40004,
              "Op": "Add", "Value": 1, "SelectorBaseSkillId": 10020101,
              "SelectorProjectileTemplateId": 30020101}),
    (600001, {"TargetKind": "SkillParameter", "TargetId": 50001,
              "Op": "Add", "Value": 1, "SelectorBaseSkillId": 10020201,
              "SelectorAreaTemplateId": 40020201}),
    (700001, {"TargetKind": "SkillParameter", "TargetId": 60003,
              "Op": "Mul", "Value": 1.5, "SelectorSummonTemplateId": 1}),
    (800001, {"TargetKind": "SkillParameter", "TargetId": 70001,
              "Op": "Add", "Value": 10, "SelectorBaseSkillId": 10020101,
              "SelectorProjectileTemplateId": 30020101}),
)


def find_sheet(workbook, field):
    matches = [sheet for sheet in workbook if field in [cell.value for cell in sheet[1]]]
    if len(matches) != 1:
        raise ValueError(f"Expected one sheet with {field}")
    return matches[0]


def headers(sheet):
    return {cell.value: cell.column for cell in sheet[1] if cell.value}


def add_row(sheet, values, prototype_row=None):
    columns = headers(sheet)
    identity = values["Id"]
    for row in range(4, sheet.max_row + 1):
        if sheet.cell(row, columns["Id"]).value != identity:
            continue
        for field, value in values.items():
            if sheet.cell(row, columns[field]).value != value:
                raise ValueError(f"Existing row {identity} differs at {field}")
        return
    row = sheet.max_row + 1
    if prototype_row is not None:
        for column in range(1, sheet.max_column + 1):
            source = sheet.cell(prototype_row, column)
            target = sheet.cell(row, column)
            target.value = source.value
            if source.has_style:
                target._style = copy(source._style)
    for field, value in values.items():
        sheet.cell(row, columns[field], value)


def ensure_column(sheet, field, kind):
    columns = headers(sheet)
    if field in columns:
        if sheet.cell(2, columns[field]).value != kind:
            raise ValueError(f"Unexpected {field} type")
        return columns[field]
    column = sheet.max_column + 1
    sheet.cell(1, column, field)
    sheet.cell(2, column, kind)
    sheet.cell(3, column, "Modifier authoring reference")
    return column


def update_effects():
    workbook = load_workbook(EFFECTS)
    templates = find_sheet(workbook, "CandidateSkillId")
    for identity, changes in TEMPLATES:
        add_row(templates, {"Id": identity, "Kind": "Buff", "Priority": 20,
                            **changes}, 4)
    buffs = find_sheet(workbook, "MaxStacks")
    ids = ",".join(str(identity) for identity, _ in TEMPLATES)
    add_row(buffs, {"Id": 102, "Name": "Modifier domain showcase",
                    "DurationMs": 5000, "IntervalMs": 0,
                    "ModifierIds": ids}, 5)
    workbook.save(EFFECTS)


def update_summons():
    workbook = load_workbook(SUMMONS)
    summons = find_sheet(workbook, "AttributeTemplateId")
    reference_column = ensure_column(summons, "InheritAttributeConfigId", "int")
    if summons.cell(4, reference_column).value not in (None, 1):
        raise ValueError("Summon 1 already has another inheritance definition")
    summons.cell(4, reference_column, 1)

    matches = [sheet for sheet in workbook
               if set(headers(sheet)) == {"##var", "Id", "OmittedFields", "NullFields"}]
    if len(matches) != 1:
        raise ValueError("Expected one summon inheritance sheet")
    inherits = matches[0]
    columns = headers(inherits)
    inherits.cell(2, columns["Id"], "int")
    ensure_column(inherits, "Scales", "string")
    scales = json.dumps([{"AttrId": 4, "Ratio": 0.25, "Add": 5.0}],
                        separators=(",", ":"))
    add_row(inherits, {"Id": 1, "Scales": scales})
    workbook.save(SUMMONS)


if __name__ == "__main__":
    update_effects()
    update_summons()
