"""Add hero-specific MOBA presentation resources and aim templates.

The workbook rows are authoritative; export_pipeline.py publishes the runtime JSON.
The placeholder prefabs are created separately under Resources/effect so every VFX
ID resolves in Unity EditMode tests.
"""

from __future__ import annotations

import copy
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[3]
DATA = ROOT / "LubanConfig/Moba/Production/Datas"


def rows(ws):
    return {row[1].value: row for row in ws.iter_rows(min_row=5, values_only=False)}


def append_row(ws, values):
    existing = rows(ws)
    if values[1] in existing:
        row = existing[values[1]]
        for index, value in enumerate(values):
            row[index].value = value
        return
    ws.append(values)
    target = list(ws.iter_rows(min_row=ws.max_row, max_row=ws.max_row))[0]
    for source in ws.iter_rows(min_row=5, max_row=5):
        for src, dst in zip(source, target):
            if src.has_style:
                dst._style = copy.copy(src._style)
                dst.number_format = src.number_format
                dst.alignment = copy.copy(src.alignment)
        break


def update_skill_buttons():
    path = DATA / "J-技能表.xlsx"
    book = openpyxl.load_workbook(path)
    ws = book[book.sheetnames[0]]
    for skill_id, template_id in {
        10070101: 11, 10070201: 12, 10070301: 13,
        10080101: 14, 10080201: 15, 10080301: 16,
    }.items():
        row = rows(ws).get(skill_id)
        if row:
            row[11].value = template_id
    book.save(path)


def add_aim_templates():
    path = DATA / "J-技能表现表.xlsx"
    book = openpyxl.load_workbook(path)
    ws = book.active
    templates = [
        (11, 450, "Direction", 12, True, True, "DashLine", 1.8, 0.35, "sun_shangxiang_roll_dash", 0, None, "Aim"),
        (12, 220, "Point", 12, True, False, "TargetCircle", 4.0, 0.35, "sun_shangxiang_grenade_area", 0, None, "TargetPoint"),
        (13, 220, "Direction", 12, True, True, "LockProjectile", 1.2, 0.35, "sun_shangxiang_cannon_line", 0, None, "Aim"),
        (14, 220, "Point", 12, True, False, "TargetCircle", 3.5, 0.35, "luban_grenade_area", 0, None, "TargetPoint"),
        (15, 220, "Direction", 12, True, True, "DirectionLine", 1.2, 0.35, "luban_shark_cannon_line", 0, None, "Aim"),
        (16, 220, "Point", 12, True, False, "TargetCircle", 4.5, 0.35, "luban_air_support_area", 0, None, "TargetPoint"),
    ]
    for item in templates:
        template_id, radius, aim_mode, threshold, enable, face, shape, width, long_press, name, select_range, self_radius, point_mode = item
        values = [None, template_id, float(radius), aim_mode, float(threshold), bool(enable), bool(face), shape, float(width), float(long_press), name, float(select_range), self_radius, point_mode, None, None]
        append_row(ws, values)
    book.save(path)


def add_vfx_rows():
    path = DATA / "T-特效表.xlsx"
    book = openpyxl.load_workbook(path)
    ws = book[book.sheetnames[0]]
    entries = [
        (90007001, "effect/sunshangxiang_skill1_roll", 600),
        (90007002, "effect/sunshangxiang_skill2_grenade", 700),
        (90007003, "effect/sunshangxiang_skill2_explosion", 1200),
        (90007004, "effect/sunshangxiang_skill3_cannon", 900),
        (90007005, "effect/sunshangxiang_buff_enhanced", 2500),
        (90007006, "effect/sunshangxiang_buff_mark", 2500),
        (90008001, "effect/luban_basic_bullet", 0),
        (90008002, "effect/luban_skill1_grenade", 700),
        (90008003, "effect/luban_skill1_explosion", 700),
        (90008004, "effect/luban_skill2_cannon", 900),
        (90008005, "effect/luban_passive_sweep", 700),
        (90008006, "effect/luban_skill3_bomb", 700),
    ]
    for entry in entries:
        append_row(ws, [None, *entry])
    templates = book[book.sheetnames[1]]
    for vfx_id, duration in ((90007001, 600), (90007002, 700), (90007003, 1200),
                             (90007004, 900), (90007005, 2500), (90007006, 2500),
                             (90008001, 0), (90008002, 700), (90008003, 700),
                             (90008004, 900), (90008005, 700), (90008006, 700)):
        values = [None, vfx_id, vfx_id, 1, 0.85, 1.0, 0.75, 1.0, duration,
                  True, 2, "placeholder_" + str(vfx_id), 0, 0.15, 0, 0, 1, None, 0, 0, None, None]
        append_row(templates, values)
    book.save(path)


def update_runtime_references():
    path = DATA / "J-投射物表.xlsx"
    book = openpyxl.load_workbook(path)
    ws = book["投射物"]
    vfx_by_projectile = {
        30070202: 90007002, 30070302: 90007004,
        30080001: 90008001, 30080101: 90008002, 30080201: 90008004,
        30080301: 90008005, 30080401: 90008006,
    }
    for projectile_id, vfx_id in vfx_by_projectile.items():
        row = rows(ws).get(projectile_id)
        if row:
            row[25].value = vfx_id
    ws = book["区域"]
    for area_id, vfx_id in {40070203: 90007003, 40080101: 90008003}.items():
        row = rows(ws).get(area_id)
        if row:
            row[17].value = vfx_id
    book.save(path)

    path = DATA / "J-效果表.xlsx"
    book = openpyxl.load_workbook(path)
    ws = book["Buff"]
    for buff_id, vfx_id in {10070101: 90007005, 10070301: 90007006}.items():
        row = rows(ws).get(buff_id)
        if row:
            row[14].value = vfx_id
    book.save(path)


def update_trigger_sources():
    import json

    path = ROOT / "Unity/Packages/com.abilitykit.demo.moba.view.runtime/Configs/MobaTriggerAuthoring/Sources/moba-passives.trigger.json"
    document = json.loads(path.read_text(encoding="utf-8"))
    for trigger in document["module"]["triggers"]:
        if trigger.get("id") != 10080002:
            continue
        action_nodes = [trigger.get("actions", {})]
        action_nodes.extend(trigger.get("actions", {}).get("children", []))
        for action in action_nodes:
            if action.get("type") != "advance_gameplay_counter":
                continue
            for argument in action.get("arguments", []):
                if argument.get("name") == "threshold":
                    argument["value"]["numberValue"] = 5.0
        trigger["note"] = "鲁班七号被动：第五次普攻触发扫射。"
    path.write_text(json.dumps(document, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main():
    update_skill_buttons()
    add_aim_templates()
    add_vfx_rows()
    update_runtime_references()
    update_trigger_sources()


if __name__ == "__main__":
    main()
