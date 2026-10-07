"""One-time migration from per-skill and numeric tags to an optional tag catalog."""

from copy import copy
from pathlib import Path

from openpyxl import load_workbook


DATA = Path(__file__).resolve().parents[1] / "Production/Datas"
PLANNED_TAGS = (
    (1100, "State.Casting", "施法过程", 1000),
    (1101, "State.Casting.Channeling", "引导中", 1100),
    (1102, "State.Casting.Charging", "蓄力中", 1100),
    (1120, "State.Movement", "运动状态", 1000),
    (1121, "State.Movement.Dashing", "冲刺中", 1120),
    (1122, "State.Movement.Leaping", "跳跃中", 1120),
    (1123, "State.Movement.Forced", "被迫位移中", 1120),
    (1140, "State.Visibility", "可见性状态", 1000),
    (1141, "State.Visibility.Invisible", "隐身", 1140),
    (1142, "State.Visibility.Camouflaged", "伪装", 1140),
    (1143, "State.Visibility.Revealed", "显形", 1140),
    (1160, "State.Defense", "防护状态", 1000),
    (1161, "State.Defense.Shielded", "护盾保护中", 1160),
    (1162, "State.Defense.DamageReduction", "减伤中", 1160),
    (1163, "State.Defense.Reflecting", "反伤中", 1160),
    (1180, "State.Attack", "攻击强化状态", 1000),
    (1181, "State.Attack.EmpoweredBasic", "强化普攻待触发", 1180),
    (1182, "State.Attack.EmpoweredSkill", "强化技能待触发", 1180),
    (1200, "State.Immunity", "免疫状态", 1000),
    (1201, "State.Immunity.Projectile", "投射物免疫", 1200),
    (1202, "State.Immunity.Slow", "减速免疫", 1200),
    (1203, "State.Immunity.Displacement", "强制位移免疫", 1200),
    (1220, "State.Survival", "生存状态", 1000),
    (1221, "State.Survival.DeathPrevented", "免死保护中", 1220),
    (1222, "State.Survival.Reviving", "复活中", 1220),
    (1500, "Debuff.Control", "控制类负面效果", 1040),
    (1501, "Debuff.Control.Taunted", "嘲讽", 1500),
    (1502, "Debuff.Control.Disarmed", "缴械", 1500),
    (1503, "Debuff.Control.Blinded", "致盲", 1500),
    (1504, "Debuff.Control.Airborne", "击飞", 1500),
    (1505, "Debuff.Control.KnockedBack", "击退", 1500),
    (1506, "Debuff.Control.Pulled", "牵引", 1500),
    (1507, "Debuff.Control.Frozen", "冰冻", 1500),
    (1508, "Debuff.Control.Petrified", "石化", 1500),
    (1520, "Debuff.Movement", "位移限制类负面效果", 1040),
    (1521, "Debuff.Movement.Grounded", "禁止位移", 1520),
    (1540, "Debuff.DamageOverTime", "持续伤害类负面效果", 1040),
    (1541, "Debuff.DamageOverTime.Burning", "灼烧", 1540),
    (1542, "Debuff.DamageOverTime.Poisoned", "中毒", 1540),
    (1543, "Debuff.DamageOverTime.Bleeding", "流血", 1540),
    (1560, "Debuff.Vision", "视野限制类负面效果", 1040),
    (1561, "Debuff.Vision.Nearsighted", "视野受限", 1560),
    (1580, "Debuff.Healing", "治疗限制类负面效果", 1040),
    (1581, "Debuff.Healing.Reduced", "治疗降低", 1580),
    (1582, "Debuff.Healing.Blocked", "禁止治疗", 1580),
    (1600, "Debuff.Defense", "防护削弱类负面效果", 1040),
    (1601, "Debuff.Defense.Vulnerable", "易伤", 1600),
    (1602, "Debuff.Defense.ArmorBroken", "物理防御削弱", 1600),
    (1603, "Debuff.Defense.MagicResistanceBroken", "法术防御削弱", 1600),
)


def replace_rows(sheet, rows):
    styles = [copy(sheet.cell(4, column)._style) for column in range(1, sheet.max_column + 1)]
    sheet.delete_rows(4, sheet.max_row - 3)
    for number, values in enumerate(rows, 4):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(number, column, value)
            cell._style = styles[column - 1]
    sheet.auto_filter.ref = f"B1:H{sheet.max_row}"


def main():
    tag_book = load_workbook(DATA / "J-标签表.xlsx")
    definitions = tag_book["标签定义"]
    for column, name, kind, comment in (
        (7, "Kind", "MobaTagKind", "Category 是层级节点，不可直接授予；Fact 才能由持续对象持有。"),
        (8, "Status", "MobaTagStatus", "Active 已接入；Planned 仅预置，引用前须先实现授予和消费。"),
    ):
        definitions.cell(1, column, name)
        definitions.cell(2, column, kind)
        definitions.cell(3, column, comment)
        for row in (1, 2, 3, 4):
            definitions.cell(row, column)._style = copy(definitions.cell(row, 4)._style)
    definitions.column_dimensions["G"].width = 18
    definitions.column_dimensions["H"].width = 18
    existing = [row for row in list(definitions.values)[3:] if isinstance(row[1], int)]
    obsolete_skill_ids = {row[1] for row in existing if row[3] == "SkillSelection" and row[2] != "Skill"}
    rows = [row for row in existing if row[3] != "LegacyNumeric" and row[1] not in obsolete_skill_ids]
    existing_ids = {row[1] for row in rows}
    for identity, name, display, parent in PLANNED_TAGS:
        if identity not in existing_ids:
            rows.append((None, identity, name, "GameplayState", display, parent, None, None))
    planned_ids = {entry[0] for entry in PLANNED_TAGS}
    category_ids = {90, 1000, 1040, *(entry[3] for entry in PLANNED_TAGS)}
    rows = [tuple(row[:6]) + (row[6] or ("Category" if row[1] in category_ids else "Fact"),
                              row[7] or ("Planned" if row[1] in planned_ids else "Active"))
            for row in rows]
    replace_rows(definitions, sorted(rows, key=lambda row: row[1]))
    definitions.cell(3, 2, "稳定 ID：90 是可选技能标签根；1000-1099 是现有战斗标签冻结段；1100-1499 为状态，1500-1999 为负面效果；2000-2999 为按需技能标签。")
    definitions.cell(3, 4, "GameplayState 为持续对象授予的战斗事实；SkillSelection 仅用于技能选择，按需配置。")

    skill_book = load_workbook(DATA / "J-技能表.xlsx")
    skills = skill_book["技能"]
    tag_column = next(column for column in range(1, skills.max_column + 1)
                      if skills.cell(1, column).value == "Tags")
    for row in range(4, skills.max_row + 1):
        cell = skills.cell(row, tag_column)
        if cell.value:
            ids = {int(value) for value in str(cell.value).split(",")}
            if not ids <= obsolete_skill_ids:
                raise ValueError(f"Skill row {row} contains a custom tag: {ids}")
            cell.value = None
    skills.cell(3, tag_column, "可选技能选择标签 ID；仅在 AI 或配置查询确有需求时填写，不按英雄或技能槽位自动分配。")

    for filename, sheet_name in (("J-召唤物表.xlsx", "召唤物"),
                                 ("J-战斗场景表.xlsx", "玩法")):
        book = load_workbook(DATA / filename)
        sheet = book[sheet_name]
        column = next(index for index in range(1, sheet.max_column + 1)
                      if sheet.cell(1, index).value == "Tags")
        if any(sheet.cell(row, column).value for row in range(4, sheet.max_row + 1)):
            raise ValueError(f"{filename}: Tags is not empty")
        sheet.cell(3, column, "暂不使用；单位身份与玩法类型由专用字段表示，不通过旧数字标签配置。")
        book.save(DATA / filename)

    tag_book.save(DATA / "J-标签表.xlsx")
    skill_book.save(DATA / "J-技能表.xlsx")
    print(f"removed {len(obsolete_skill_ids)} per-skill tags; catalog now has {len(rows)} definitions")


if __name__ == "__main__":
    main()
