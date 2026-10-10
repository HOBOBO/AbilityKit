"""One-time migration of simple control conditions to flat tag columns."""

from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path(__file__).resolve().parents[1] / "Production/Datas/J-标签表.xlsx"
CONTROL_TEMPLATES = {10000001, *range(3005, 3012)}
CONTROL_IMMUNITY_TAGS = "1001,1002"


def migrate():
    book = load_workbook(WORKBOOK)
    sheet = book["持续标签模板"]
    columns = {cell.value: cell.column for cell in sheet[1] if cell.value}
    rows = {sheet.cell(index, columns["Id"]).value: index
            for index in range(4, sheet.max_row + 1)}
    if not CONTROL_TEMPLATES <= rows.keys():
        raise ValueError(f"missing control templates: {CONTROL_TEMPLATES - rows.keys()}")

    for template_id in sorted(CONTROL_TEMPLATES):
        row = rows[template_id]
        activation_tags = sheet.cell(row, columns["ActivationBlockedTagIds"])
        activation_rule = sheet.cell(row, columns["ActivationRuleGroupId"])
        if activation_tags.value not in (None, CONTROL_IMMUNITY_TAGS) or activation_rule.value not in (0, 10000001):
            raise ValueError(f"template {template_id}: unexpected activation condition")
        activation_tags.value = CONTROL_IMMUNITY_TAGS
        activation_rule.value = 0

        if template_id == 10000001:
            ongoing_tags = sheet.cell(row, columns["OngoingBlockedTagIds"])
            ongoing_rule = sheet.cell(row, columns["OngoingRuleGroupId"])
            if ongoing_tags.value not in (None, CONTROL_IMMUNITY_TAGS) or ongoing_rule.value not in (0, 10000001):
                raise ValueError("forced motion template: unexpected ongoing condition")
            ongoing_tags.value = CONTROL_IMMUNITY_TAGS
            ongoing_rule.value = 0

    notes = {
        "ActivationRequiredTagIds": "激活前必须全部具备的标签 ID；与激活规则组同时配置时两者都要满足。空表示无此要求。",
        "ActivationBlockedTagIds": "激活前任一命中即阻断的标签 ID；简单 OR 条件优先填此列。空表示无此阻断。",
        "ApplicationTagIds": "持续对象处于生效状态时授予的事实标签 ID；暂停或结束后不再授予。空表示只配置门禁。",
        "RemovalRequiredTagIds": "持续期间全部具备时触发移除的标签 ID；与移除规则组同时配置时两者都要满足。",
        "RemovalBlockedTagIds": "持续期间这些标签全部不存在时才满足移除条件；不是“遇到这些标签就移除”。",
        "OngoingRequiredTagIds": "持续期间必须全部具备的标签 ID；不满足时暂停，恢复后继续。空表示无此要求。",
        "OngoingBlockedTagIds": "持续期间任一命中即暂停的标签 ID；标签消失后可恢复。空表示无此阻断。",
        "ActivationRuleGroupId": "复合激活条件规则组 ID；与平铺条件按 AND 合并。简单 Required/Blocked 优先使用标签列；0 表示不使用。",
        "RemovalRuleGroupId": "复合移除条件规则组 ID；与平铺条件按 AND 合并。整组留空时不会自动移除；0 表示不使用。",
        "OngoingRuleGroupId": "复合持续条件规则组 ID；与平铺条件按 AND 合并。不满足时暂停而非移除；0 表示不使用。",
    }
    for field, note in notes.items():
        sheet.cell(3, columns[field]).value = note
    book.save(WORKBOOK)


if __name__ == "__main__":
    migrate()
