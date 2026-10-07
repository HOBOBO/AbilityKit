"""One-time removal of retired MOBA workbook tables and trigger snapshots."""

from pathlib import Path

from openpyxl import load_workbook


DATAS = Path(__file__).resolve().parents[1] / "Production/Datas"
RETIRED = {
    "J-标签表.xlsx": ("标签模板", "TagTemplates"),
    "J-角色表.xlsx": ("旧属性模板", "ResourceAttributetemplates"),
    "J-资源节点表.xlsx": ("演示物品", "ResourceDemoTbitem"),
}


def main():
    catalog_path = DATAS / "__tables__.xlsx"
    catalog = load_workbook(catalog_path)
    catalog_sheet = catalog.active
    names = {row[1]: index for index, row in enumerate(catalog_sheet.values, 1)
             if isinstance(row[1], str)}
    for filename, (sheet_name, table_name) in RETIRED.items():
        path = DATAS / filename
        workbook = load_workbook(path)
        if sheet_name not in workbook or table_name not in names:
            raise ValueError(f"missing retired table {table_name} in {filename}")
        del workbook[sheet_name]
        if filename == "J-资源节点表.xlsx":
            nodes = workbook["Ability节点"]
            for index in range(nodes.max_row, 3, -1):
                path_value = nodes.cell(index, 4).value
                if isinstance(path_value, str) and path_value.startswith("ability/triggers/"):
                    nodes.delete_rows(index)
        workbook.save(path)
    for index in sorted((names[table_name] for _, table_name in RETIRED.values()), reverse=True):
        catalog_sheet.delete_rows(index)
    catalog.save(catalog_path)


if __name__ == "__main__":
    main()
