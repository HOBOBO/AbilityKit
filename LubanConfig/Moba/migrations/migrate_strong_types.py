"""One-time migration of the production MOBA workbooks to Luban beans/enums."""

import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DATA_DIR = Path(__file__).resolve().parents[1] / "Production/Datas"

ENUMS = {
    "BuffStackingPolicy": ["None", "IgnoreIfExists", "Replace", "AddStack", "RefreshDuration"],
    "BuffRefreshPolicy": ["None", "KeepRemaining", "ResetRemaining", "AddRemaining"],
    "BuffDispelPolicy": ["LegacyTag", "Dispellable", "Undispellable"],
    "MobaComponentOpKind": ["None", "SetModelId", "SetLifetimeMs"],
    "MobaModifierTargetKind": {1: "Attribute", 2: "StateMachineParameter", 3: "StateFlag",
                              4: "SkillParameter", 255: "Custom"},
    "ModifierOp": {0: "Add", 1: "Mul", 2: "Override", 3: "PercentAdd", 100: "Custom"},
    "MagnitudeSourceType": ["Fixed", "Scalable", "Attribute", "TimeDecay", "Pipeline", "ContextFloat"],
    "ModifierEvaluationPolicy": ["Realtime", "OnApplySnapshot"],
    "SearchQueryExplicitTargetPolicy": ["PreferExplicitTarget", "SearchWhenMissing", "IgnoreExplicitTarget"],
    "SearchTargetProviderKind": {1: "AllActors", 2: "ExplicitTarget", 3: "ContextTarget", 4: "Caster",
                                 5: "SameTeam", 6: "EnemyTeam", 7: "MainType", 8: "UnitSubType"},
    "SearchTargetSelectorKind": {4097: "TopKByScore", 4098: "StreamingTopKByScore"},
    "SearchTargetRuleKind": {257: "CircleShape", 258: "SectorShape", 513: "Whitelist",
                             514: "Blacklist", 516: "RequireValidId", 517: "RequireHasPosition",
                             769: "ExcludeCaster", 770: "ExcludeExplicitTarget",
                             771: "RectangleShape", 772: "CapsuleShape"},
    "SearchTargetScorerKind": {8193: "Zero", 8194: "SeededHashRandom",
                               8196: "DistanceToCaster", 8197: "DistanceToExplicitTarget"},
    "SearchTargetPointKind": ["Caster", "AimPosition", "ExplicitTarget"],
    "MapCollisionShapeType": ["Box", "Sphere", "Capsule"],
    "MapProjectileResponse": ["Default", "Block", "Ignore", "HitOnly", "Pierce"],
    "BattleRunMode": ["Normal", "Record", "Replay"],
    "BattleSyncMode": ["Lockstep", "SnapshotAuthority", "StateSync", "Hybrid"],
    "SkillType": ["None", "NormalAttack", "Active", "Passive", "Ultimate", "ParallelActive"],
    "SkillAimMode": ["Direction", "Point"],
    "SkillAimIndicatorShape": ["Hidden", "DirectionLine", "TargetCircle", "SelfCircle", "Sector",
                               "DirectionArea", "DashLine", "LockProjectile", "FanArea"],
    "SkillUsePointMode": ["None", "Aim", "TargetPoint"],
    "ProjectileEmitterType": ["None", "Linear"],
    "ProjectileHitPolicyKind": ["ExitOnHit", "Pierce"],
    "SearchQueryCenterMode": ["Caster", "AimPos", "ExplicitTarget"],
}

LEVEL = [("CooldownMs", "int"), ("Cost", "int"), ("Params", "(list#sep=,),float")]
OP = [("Kind", "MobaComponentOpKind"), ("IntValue", "int"),
      ("FloatValue", "float"), ("BoolValue", "bool")]
MODIFIER = [
    ("TargetKind", "MobaModifierTargetKind"), ("TargetId", "int"), ("AttrTypeId", "int"),
    ("Op", "ModifierOp"), ("Value", "float"), ("Priority", "int"),
    ("MagnitudeSourceType", "MagnitudeSourceType"), ("MagnitudeBaseValue", "float"),
    ("MagnitudeCoefficient", "float"),
]
BUFF_MODIFIER = MODIFIER + [
    ("MagnitudeDuration", "float"), ("MagnitudeDecayType", "int"),
    ("MagnitudeAttributeTypeId", "int"), ("MagnitudeContextKey", "string"),
    ("MagnitudeCurve", "(list#sep=,),float"), ("EvaluationPolicy", "ModifierEvaluationPolicy"),
]
PROCESS_MODIFIER = MODIFIER + [("MagnitudeContextKey", "string")]
VECTOR = "MapVector3"

# Header signature, expanded columns, scalar enum columns.
MIGRATIONS = [
    ({"Levels"}, {"Levels": ("SkillLevel", LEVEL, True)}, {}),
    ({"Ops"}, {"Ops": ("ComponentOp", OP, True)}, {}),
    ({"Modifiers", "DispelPolicy"}, {"Modifiers": ("BuffModifier", BUFF_MODIFIER, True)},
     {"DispelPolicy": "BuffDispelPolicy", "RefreshPolicy": "BuffRefreshPolicy",
      "StackingPolicy": "BuffStackingPolicy"}),
    ({"Modifiers", "RequireOutOfCombat"},
     {"Modifiers": ("ContinuousProcessModifier", PROCESS_MODIFIER, True)}, {}),
    ({"Provider", "Selector", "ExplicitTargetPolicy"},
     {"Provider": ("SearchTargetProvider", [("Kind", "SearchTargetProviderKind")], False),
      "Selector": ("SearchTargetSelector", [("Kind", "SearchTargetSelectorKind")], False)},
     {"ExplicitTargetPolicy": "SearchQueryExplicitTargetPolicy"}),
    ({"Rules", "Scorers", "Provider"},
     {"Rules": ("SearchTargetRule", [("Kind", "SearchTargetRuleKind"),
                                     ("Center", "SearchTargetPointKind?"),
                                     ("Forward", "SearchTargetPointKind?"),
                                     ("Radius", "float?"), ("HalfAngleDeg", "float?")], True),
      "Scorers": ("SearchTargetScorer", [("Kind", "SearchTargetScorerKind"),
                                         ("Source", "SearchTargetPointKind?"),
                                         ("RandomSeed", "int?")], True)}, {}),
    ({"Bounds", "SpawnPoints", "CollisionObjects"},
     {"Bounds": ("MapBounds", [("Center", VECTOR), ("Size", VECTOR)], False),
      "WalkableAreas": ("MapWalkableArea", [("Id", "int"), ("Name", "string"),
                                              ("Center", VECTOR), ("Size", VECTOR)], True),
      "SpawnPoints": ("MapSpawnPoint", [("Id", "int"), ("TeamId", "int"),
                                          ("Position", VECTOR), ("YawDegrees", "float")], True),
      "CollisionObjects": ("MapCollisionObject", [
          ("Id", "int"), ("Name", "string"), ("ObjectType", "string"),
          ("ShapeType", "MapCollisionShapeType"),
          ("Position", VECTOR), ("RotationEuler", VECTOR), ("Size", VECTOR), ("Radius", "float"),
          ("Height", "float"), ("CollisionLayer", "int"), ("BlocksMovement", "bool"),
          ("ProjectileResponse", "MapProjectileResponse"), ("GenerateView", "bool")], True)}, {}),
    ({"Network", "Players", "RunMode"},
     {"Network": ("BattleNetwork", [("Host", "string"), ("Port", "int")], False),
      "Players": ("BattlePlayer", [
          ("PlayerId", "string"), ("TeamId", "int"), ("HeroId", "int"), ("Name", "string"),
          ("PositionX", "float"), ("PositionZ", "float"), ("BrainId", "int?"),
          ("EnableBrainOnSpawn", "bool?")], True)}, {}),
]
BEAN_NAMES = {bean_name for _, groups, _ in MIGRATIONS
              for bean_name, _, _ in groups.values()}
SCALAR_ENUM_MIGRATIONS = [
    ({"SkillType", "CastFlowId"}, {"SkillType": "SkillType"}),
    ({"AimMode", "IndicatorShape", "UsePointMode"},
     {"AimMode": "SkillAimMode", "IndicatorShape": "SkillAimIndicatorShape",
      "UsePointMode": "SkillUsePointMode"}),
    ({"EmitterType", "CountPerShot"}, {"EmitterType": "ProjectileEmitterType"}),
    ({"HitPolicyKind", "HitsRemaining"}, {"HitPolicyKind": "ProjectileHitPolicyKind"}),
    ({"CenterMode", "EmitKind"}, {"CenterMode": "SearchQueryCenterMode"}),
]


def enum_name(value, type_name):
    if value is None or isinstance(value, str):
        return value
    values = ENUMS[type_name]
    if isinstance(values, list):
        if value < 0 or value >= len(values):
            raise ValueError(f"Unknown {type_name}: {value}")
        return values[value]
    return values[value]


def cell_value(value, type_name):
    if type_name == VECTOR and value is not None:
        return ",".join(str(value[axis]) for axis in ("X", "Y", "Z"))
    enum_type = type_name.removesuffix("?")
    if enum_type in ENUMS:
        return enum_name(value, enum_type)
    if isinstance(value, list):
        return ",".join(str(item) for item in value)
    return value


def migrate_sheet(book, source, groups, scalar_enums):
    index = book.sheetnames.index(source.title)
    original_headers = [cell.value for cell in source[1]]
    original_types = [cell.value for cell in source[2]]
    original_members = [cell.value for cell in source[3]]
    source_data = []
    for row in source.iter_rows(min_row=4, values_only=True):
        if any(value is not None for value in row):
            source_data.append(dict(zip(original_headers, row)))

    title = source.title
    book.remove(source)
    sheet = book.create_sheet(title, index)
    headers, types, members = [], [], []
    for name, type_name, member in zip(original_headers, original_types, original_members):
        if name in groups:
            bean_name, fields, is_list = groups[name]
            headers.extend([("*" if is_list else "") + name] + [None] * (len(fields) - 1))
            types.extend([("list," if is_list else "") + bean_name] + [None] * (len(fields) - 1))
            members.extend([field for field, _ in fields])
        else:
            headers.append(name)
            types.append(scalar_enums.get(name, type_name))
            members.append(member)
    sheet.append(headers)
    sheet.append(types)
    sheet.append(["##var"] + members[1:])

    offset = 1
    for name in original_headers:
        if name not in groups:
            offset += 1
            continue
        size = len(groups[name][1])
        if size > 1:
            sheet.merge_cells(start_row=1, start_column=offset, end_row=1, end_column=offset + size - 1)
            sheet.merge_cells(start_row=2, start_column=offset, end_row=2, end_column=offset + size - 1)
        offset += size

    for source_row in source_data:
        parsed = {}
        count = 1
        for name, (_, _, is_list) in groups.items():
            raw = source_row[name]
            value = json.loads(raw) if isinstance(raw, str) and raw else raw
            parsed[name] = value
            if is_list and value:
                count = max(count, len(value))
        for item_index in range(count):
            output = []
            for name in original_headers:
                if name not in groups:
                    value = source_row[name] if item_index == 0 else None
                    output.append(enum_name(value, scalar_enums[name]) if name in scalar_enums else value)
                    continue
                _, fields, is_list = groups[name]
                value = parsed[name]
                item = value[item_index] if is_list and value and item_index < len(value) else (
                    value if not is_list and item_index == 0 else None)
                output.extend(cell_value(item.get(field), type_name) if item else None
                              for field, type_name in fields)
            sheet.append(output)

    blue = PatternFill("solid", fgColor="DCEAF3")
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10, bold=cell.row <= 3)
            cell.alignment = Alignment(vertical="center")
            if cell.row <= 3:
                cell.fill = blue
    for index, column in enumerate(sheet.columns, 1):
        field = column[2].value if len(column) > 2 else None
        sheet.column_dimensions[get_column_letter(index)].width = max(14, min(30, len(str(field or "")) + 3))
    sheet.freeze_panes = "C4"
    sheet.auto_filter.ref = f"A3:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
    return title


def main():
    remaining = list(MIGRATIONS)
    for path in sorted(DATA_DIR.glob("[JMT]-*.xlsx")):
        book = load_workbook(path)
        changed = []
        typed_sheets = set()
        for sheet in list(book):
            headers = {cell.value.lstrip("*") if isinstance(cell.value, str) else cell.value
                       for cell in sheet[1]}
            for migration in list(remaining):
                signature, groups, enums = migration
                if signature <= headers:
                    typed_sheets.add(sheet.title)
                    already_typed = all(
                        sheet.cell(2, next(i for i, cell in enumerate(sheet[1], 1)
                                           if cell.value in (name, "*" + name))).value
                        == ("list," if is_list else "") + bean_name
                        for name, (bean_name, _, is_list) in groups.items()
                    )
                    if not already_typed:
                        changed.append(migrate_sheet(book, sheet, groups, enums))
                        sheet = book[sheet.title]
                        headers = {cell.value.lstrip("*") if isinstance(cell.value, str) else cell.value
                                   for cell in sheet[1]}
                    remaining.remove(migration)
        for sheet in book:
            if sheet.title in typed_sheets:
                for column in range(2, sheet.max_column + 1):
                    type_name = sheet.cell(2, column).value
                    member = sheet.cell(3, column)
                    if type_name and type_name.removeprefix("list,") not in BEAN_NAMES and member.value:
                        member.value = None
                        if sheet.title not in changed:
                            changed.append(sheet.title)
            headers = [cell.value for cell in sheet[1]]
            if "RunMode" in headers and "SyncMode" in headers and "*Players" in headers:
                for name, enum_type in (("RunMode", "BattleRunMode"), ("SyncMode", "BattleSyncMode")):
                    column = headers.index(name) + 1
                    if sheet.cell(2, column).value != enum_type:
                        sheet.cell(2, column).value = enum_type
                        if sheet.title not in changed:
                            changed.append(sheet.title)
            normalized_headers = {name.lstrip("*") if isinstance(name, str) else name
                                  for name in headers}
            for signature, fields in SCALAR_ENUM_MIGRATIONS:
                if not signature <= normalized_headers:
                    continue
                for name, enum_type in fields.items():
                    column = headers.index(name) + 1
                    if sheet.cell(2, column).value == enum_type:
                        continue
                    sheet.cell(2, column).value = enum_type
                    for row in range(4, sheet.max_row + 1):
                        cell = sheet.cell(row, column)
                        cell.value = enum_name(cell.value, enum_type)
                    if sheet.title not in changed:
                        changed.append(sheet.title)
        if changed:
            book.save(path)
            print(f"{path.name}: {', '.join(changed)}")
    if remaining:
        raise RuntimeError(f"Missing sheets for {len(remaining)} migrations")


if __name__ == "__main__":
    main()
