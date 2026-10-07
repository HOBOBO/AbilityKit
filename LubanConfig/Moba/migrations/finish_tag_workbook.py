"""Finish the initial tag workbook migration and validate its cells."""
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font

DATAS = Path(__file__).resolve().parents[1] / "Production/Datas"


def main():
    book = load_workbook(DATAS / "J-标签表.xlsx")
    descriptions = {
        "标签定义": "标签事实字典；状态标签通过 ID 引用，技能分类标签不能用于运行态规则。",
        "标签规则组": "可复用条件组；RootNodeId 指向本组唯一根节点。",
        "标签规则节点": "逐行维护条件树；Tag 使用 TagId，GroupRef 使用 RefGroupId，Not 只能有一个子节点。",
        "结果标志定义": "代码消费的结果位；已发布的 Key、BitIndex 禁止重新分配。",
        "标签结果映射": "规则组满足时设置结果位；同一结果的多个映射按 OR 合并。",
        "标签模板": "授予、移除使用标签 ID；生效条件引用规则组；授予保持来源引用计数。",
        "持续标签模板": "激活、移除、持续条件分别引用规则组；应用与移除使用标签 ID。",
    }
    display_names = {"Untargetable": "不可被选取", "Invulnerable": "免疫伤害", "CannotCast": "不能施法",
                     "Stunned": "硬控中", "CannotMove": "不能移动", "ControlImmune": "控制免疫",
                     "CannotBeControlled": "不能受控制", "CannotAttack": "不能攻击"}
    flags = book["结果标志定义"]
    if not any(row[2].value == "CannotAttack" for row in list(flags.rows)[3:]):
        flags.append([None, 8, "CannotAttack", 7, display_names["CannotAttack"]])
        book["标签结果映射"].append([None, 8, 8, 2004])
    groups = book["标签规则组"]
    nodes = book["标签规则节点"]
    tag_ids = {row[2]: row[1] for row in list(book["标签定义"].values)[3:]}
    if not any(row[1].value == 2008 for row in list(groups.rows)[3:]):
        root = max(row[1] for row in list(nodes.values)[3:]) + 1
        groups.append([None, 2008, "不能攻击：硬控、恐惧与睡眠", root])
        nodes.append([None, root, 2008, 0, 0, "Any", 0, False, 0])
        nodes.append([None, root + 1, 2008, root, 1, "GroupRef", 0, False, 2004])
        for order, name in enumerate(("State.Feared", "Feared", "feared", "State.Asleep",
                                      "Asleep", "Sleeping", "asleep", "sleeping"), 2):
            nodes.append([None, root + order, 2008, root, order, "Tag", tag_ids[name], False, 0])
    for row in list(book["标签结果映射"].rows)[3:]:
        if row[2].value == 8:
            row[3].value = 2008
    for row in list(flags.rows)[3:]:
        row[4].value = display_names[row[2].value]
    for sheet in book:
        sheet["B3"] = descriptions[sheet.title]
        sheet.auto_filter.ref = f"B4:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
        for row in sheet:
            for cell in row:
                if cell.data_type == "f":
                    raise ValueError("The tag workbook must not contain formulas.")
                cell.font = Font(name="Arial", size=10, bold=cell.row <= 2)
    book.save(DATAS / "J-标签表.xlsx")

    effects = load_workbook(DATAS / "J-效果表.xlsx")
    buffs = effects["Buff"]
    columns = {cell.value: cell.column for cell in buffs[1]}
    for row in buffs.iter_rows(min_row=4):
        if row[columns["Id"] - 1].value not in (10060101, 10060211):
            continue
        cell = row[columns["TagIds"] - 1]
        existing = str(cell.value or "").split(",")
        slow_id = str(tag_ids["Debuff.Slow"])
        cell.value = ",".join(value for value in existing if value) if slow_id in existing else ",".join(
            [value for value in existing if value] + [slow_id])
        for field in ("NullFields", "OmittedFields"):
            metadata = row[columns[field] - 1]
            metadata.value = ",".join(value for value in str(metadata.value or "").split(",")
                                      if value and value != "TagNames") or None
    effects.save(DATAS / "J-效果表.xlsx")


if __name__ == "__main__":
    main()
