"""One-time removal of graph and trigger-plan sheets from the Luban source."""

from pathlib import Path

from openpyxl import load_workbook


DATA = Path(__file__).resolve().parents[1] / "Production" / "Datas"
MOVED = {
    "actor_state_machines": ("J-角色表.xlsx", "状态机"),
    "character_hfsm": ("J-角色表.xlsx", "角色HFSM"),
    "effect_plans": ("J-效果表.xlsx", "效果计划"),
}


def main():
    registry_path = DATA / "__tables__.xlsx"
    registry = load_workbook(registry_path)
    sheet = registry.active
    outputs = {sheet.cell(row, 11).value: row for row in range(4, sheet.max_row + 1)}
    for name in MOVED:
        row = outputs.get(name)
        if row is not None:
            sheet.delete_rows(row)
            outputs = {sheet.cell(index, 11).value: index
                       for index in range(4, sheet.max_row + 1)}
    registry.save(registry_path)

    for workbook_name, sheet_name in MOVED.values():
        path = DATA / workbook_name
        workbook = load_workbook(path)
        if sheet_name in workbook:
            del workbook[sheet_name]
            workbook.save(path)


if __name__ == "__main__":
    main()
