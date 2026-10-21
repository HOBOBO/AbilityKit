"""Add the 1007 Sun Shangxiang demo slice to the production MOBA sources.

This is a one-time source migration. It copies only the structural rows needed
by the 1007 slice from existing production heroes, then assigns the 1007 ID
namespace. Generated JSON/bytes are produced by export_pipeline.py.
"""

from __future__ import annotations

import copy
import json
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[3]
DATA = ROOT / "LubanConfig/Moba/Production/Datas"
TRIGGER_ROOT = ROOT / "Unity/Packages/com.abilitykit.demo.moba.view.runtime/Configs/MobaTriggerAuthoring/Sources"


def workbook(name: str):
    return openpyxl.load_workbook(DATA / name)


def row_values(ws, key: int):
    for row in ws.iter_rows(min_row=5, values_only=False):
        if row[1].value == key:
            return row
    raise KeyError(f"{ws.title}: missing row {key}")


def deduplicate_ids(ws):
    seen = set()
    for index in range(ws.max_row, 4, -1):
        key = ws.cell(index, 2).value
        if key in seen:
            ws.delete_rows(index, 1)
        else:
            seen.add(key)


def copy_row(ws, source_id: int, target_id: int, replacements: dict[int, object] | None = None):
    if any(row[1].value == target_id for row in ws.iter_rows(min_row=5, values_only=False)):
        return
    source = row_values(ws, source_id)
    values = [cell.value for cell in source]
    values[1] = target_id
    for index, value in (replacements or {}).items():
        values[index] = value
    ws.append(values)
    target = list(ws.iter_rows(min_row=ws.max_row, max_row=ws.max_row))[0]
    for source_cell, target_cell in zip(source, target):
        if source_cell.has_style:
            target_cell._style = copy.copy(source_cell._style)
        target_cell.number_format = source_cell.number_format
        target_cell.alignment = copy.copy(source_cell.alignment)


def replace_namespace(value, old: str = "1006", new: str = "1007"):
    if isinstance(value, str):
        return value.replace(old, new)
    return value


def deduplicate_trigger_ids(triggers):
    """Keep one authored definition per trigger ID when this migration is rerun.

    The passive cooldown-reset trigger is the one case where the latest custom
    definition is authoritative; older reruns left cloned five-argument
    versions in the source list.
    """
    selected = {}
    for trigger in triggers:
        trigger_id = trigger["id"]
        if trigger_id == 10070002 and "reset_cooldown" in json.dumps(trigger, ensure_ascii=False):
            selected[trigger_id] = trigger
        elif trigger_id not in selected:
            selected[trigger_id] = trigger
    return list(selected.values())


def remap_trigger_namespace(trigger, source_namespace="1006", target_namespace="1007"):
    """Remap nested resource IDs as well as skill IDs in a trigger snapshot."""
    raw = json.dumps(trigger, ensure_ascii=False)
    for prefix in ("3106", "3006", "4006", "5006",
                   "3104", "3004", "4004",
                   "3103", "3003", "4003", "5003", source_namespace):
        raw = raw.replace(prefix, prefix[:-1] + target_namespace[-1])
    remapped = json.loads(raw)
    remapped["id"] = trigger["id"]
    return remapped


def append_excel_sources():
    # Character and attribute template.
    book = workbook("J-角色表.xlsx")
    ws = book["角色"]
    copy_row(ws, 1006, 1007, {2: 1007, 3: 1007, 4: "孙尚香", 5: "10070000", 6: "10070101,10070201,10070301"})
    ws = book["属性模板"]
    copy_row(ws, 1006, 1007, {2: 0, 3: 0, 4: 10070001})
    book.save(DATA / "J-角色表.xlsx")

    # Skill metadata and level parameters.
    book = workbook("J-技能表.xlsx")
    ws = book["技能"]
    copy_row(ws, 10060001, 10070001, {2: 10070001, 3: 7, 4: 800, 5: 701, 6: 0, 7: "孙尚香-普通攻击", 9: 8, 11: 1, 12: "NormalAttack"})
    copy_row(ws, 10060101, 10070101, {2: 10070101, 3: 7, 4: 7000, 5: 701, 6: 10070101, 7: "孙尚香-翻滚突袭", 9: 8, 11: 8, 12: "Active"})
    copy_row(ws, 10060201, 10070201, {2: 10070201, 3: 7, 4: 9000, 5: 702, 6: 10070201, 7: "孙尚香-红莲爆弹", 9: 12, 11: 2, 12: "Active"})
    copy_row(ws, 10060301, 10070301, {2: 10070301, 3: 7, 4: 20000, 5: 703, 6: 10070301, 7: "孙尚香-究极弩炮", 9: 18, 11: 3, 12: "Ultimate"})
    ws = book["技能等级"]
    copy_row(ws, 10060000, 10070000, {2: 0, 3: 0, 4: "500", 5: "basicAttackCooldownResetMs"})
    passive_level = row_values(ws, 10070000)
    passive_level[4].value = "500"
    passive_level[5].value = "basicAttackCooldownResetMs"
    copy_row(ws, 10060101, 10070101, {2: 7000, 3: 50, 4: "180,8,450,400,2500,0.35,1500", 5: "damage,dashDistance,dashDurationMs,dashSpeed,enhancedWindowMs,enhancedAttackRatio,armorReductionDurationMs"})
    copy_row(ws, 10060201, 10070201, {2: 9000, 3: 60, 4: "260,0.6,12,4,2000,0.35,1500", 5: "damage,physicalAttackRatio,range,radius,slowDurationMs,armorReductionPercent,armorReductionDurationMs"})
    copy_row(ws, 10060301, 10070301, {2: 20000, 3: 100, 4: "360,1.0,18,500,3000", 5: "damage,physicalAttackRatio,range,projectileSpeed,damageFalloffPercent"})
    ws = book.worksheets[1]
    copy_row(ws, 10060000, 10070000, {4: "孙尚香-活力迸发", 5: "10070000,10070002,10070003"})
    book.save(DATA / "J-技能表.xlsx")

    # Buffs: enhanced basic attack, slow/debuff and ultimate marker.
    book = workbook("J-效果表.xlsx")
    ws = book["Buff"]
    copy_row(ws, 10030101, 10070101, {10: "孙尚香-翻滚突袭-强化普攻窗口", 8: 1, 18: "10070111"})
    copy_row(ws, 10030102, 10070102, {10: "孙尚香-翻滚突袭-减速", 8: 1})
    copy_row(ws, 10030104, 10070104, {10: "孙尚香-翻滚突袭-破甲", 8: 1})
    copy_row(ws, 10060211, 10070201, {10: "孙尚香-红莲爆弹-减速", 8: 1})
    copy_row(ws, 10060301, 10070301, {10: "孙尚香-究极弩炮-命中标记", 8: 1})
    book.save(DATA / "J-效果表.xlsx")

    # Projectile, launcher and area rows. Existing VFX/model references are
    # intentionally reused until dedicated Sun Shangxiang assets are supplied.
    book = workbook("J-投射物表.xlsx")
    for sheet, mappings in {
        "投射发射器": [(31060001, 31070001), (31060301, 31070301), (31040201, 31070201)],
        "投射物": [(30060001, 30070001), (30060301, 30070301), (30040201, 30070201)],
        "区域": [(40060101, 40070101), (40060201, 40070201), (40040201, 40070202)],
    }.items():
        ws = book[sheet]
        for source_id, target_id in mappings:
            if any(row[1].value == target_id for row in ws.iter_rows(min_row=5, values_only=False)):
                continue
            source = row_values(ws, source_id)
            values = [replace_namespace(cell.value) for cell in source]
            values[1] = target_id
            # Keep existing state machine/VFX assets; only event/trigger IDs
            # move to the 1007 namespace.
            ws.append(values)
            target = list(ws.iter_rows(min_row=ws.max_row, max_row=ws.max_row))[0]
            for source_cell, target_cell in zip(source, target):
                if source_cell.has_style:
                    target_cell._style = copy.copy(source_cell._style)
        deduplicate_ids(ws)
    book.save(DATA / "J-投射物表.xlsx")

    # The copied projectile rows contain nested trigger references in string
    # and integer columns; normalize those references explicitly after the
    # structural copy.
    book = workbook("J-投射物表.xlsx")
    ws = book["投射物"]
    projectile_refs = {
        30070001: {11: None, 13: 10070011},
        30070201: {11: "10070211"},
        30070301: {13: 10070313},
    }
    for target_id, fields in projectile_refs.items():
        row = row_values(ws, target_id)
        for index, value in fields.items():
            row[index].value = value
    ws = book["区域"]
    for target_id, fields in {
        40070101: {13: "10070111", 15: "10070111"},
        40070201: {13: "10070211"},
        40070202: {13: "10070212", 14: None, 15: "10070212"},
    }.items():
        row = row_values(ws, target_id)
        for index, value in fields.items():
            row[index].value = value
    book.save(DATA / "J-投射物表.xlsx")


def clone_triggers():
    # Clone stable, already-published action graphs. The namespace replacement
    # also updates nested references (Buff/Area/Projectile/Skill IDs).
    skill_path = TRIGGER_ROOT / "moba-skills.trigger.json"
    skills = json.loads(skill_path.read_text(encoding="utf-8"))
    skills["module"]["triggers"] = deduplicate_trigger_ids(skills["module"]["triggers"])
    skills["module"]["triggers"] = [
        remap_trigger_namespace(trigger) if 10070000 <= trigger["id"] < 10080000 else trigger
        for trigger in skills["module"]["triggers"]
    ]
    templates = {10060001, 10060011, 10030101, 10030111,
                 10040201, 10040211, 10040212, 10060301, 10060311, 10060313}
    existing = {trigger["id"] for trigger in skills["module"]["triggers"]}
    for trigger in list(skills["module"]["triggers"]):
        if trigger["id"] not in templates:
            continue
        clone = json.loads(json.dumps(trigger))
        clone = remap_trigger_namespace(clone)
        raw = json.dumps(clone, ensure_ascii=False)
        raw = raw.replace("1003", "1007").replace("1004", "1007")
        clone = json.loads(raw)
        clone["id"] = int(str(trigger["id"]).replace("1006", "1007").replace("1003", "1007").replace("1004", "1007"))
        clone["name"] = clone["name"].replace("嬴政", "孙尚香").replace("赵云", "孙尚香").replace("墨子", "孙尚香")
        clone["groupPath"] = f"trigger_{clone['id']}"
        if clone["id"] not in existing:
            skills["module"]["triggers"].append(clone)
            existing.add(clone["id"])

    def walk(node):
        if isinstance(node, dict):
            yield node
            for value in node.values():
                yield from walk(value)
        elif isinstance(node, list):
            for value in node:
                yield from walk(value)

    def argument(node, name):
        return next(arg["value"] for arg in node.get("arguments", []) if arg["name"] == name)

    # The first skill owns the actual roll motion. Reuse the existing dash
    # action descriptor and keep the movement in the trigger layer.
    dash_source = next(t for t in skills["module"]["triggers"] if t["id"] == 10030102)
    dash = next(n for n in walk(dash_source) if n.get("type") == "dash")
    dash = json.loads(json.dumps(dash))
    dash = json.loads(json.dumps(dash).replace("100301", "100701"))
    skill_one = next(t for t in skills["module"]["triggers"] if t["id"] == 10070101)
    children = skill_one["actions"].setdefault("children", [])
    dash_children = [child for child in children if child.get("type") == "dash"]
    children[:] = [child for index, child in enumerate(children)
                   if child.get("type") != "dash" or index == next(
                       (i for i, candidate in enumerate(children)
                        if candidate.get("type") == "dash"), index)]
    if not dash_children:
        children.insert(0, dash)

    skill_two = next(t for t in skills["module"]["triggers"] if t["id"] == 10070211)
    for node in walk(skill_two):
        if node.get("type") == "spawn_area":
            argument(node, "area_id")["integerValue"] = 40070202
    skill_two_hit = next(t for t in skills["module"]["triggers"] if t["id"] == 10070212)
    for node in walk(skill_two_hit):
        for arg in node.get("arguments", []):
            if arg["name"] == "target_query_id":
                arg["value"]["integerValue"] = 50040201
    for trigger_id, damage, ratio, reason_param in (
        (10070111, 220.0, 1.0, 10070101),
        (10070212, 260.0, 0.6, 10070201),
        (10070313, 360.0, 1.0, 10070301),
    ):
        trigger = next(t for t in skills["module"]["triggers"] if t["id"] == trigger_id)
        for node in walk(trigger):
            if node.get("type") != "give_damage":
                continue
            argument(node, "damage_value")["numberValue"] = damage
            argument(node, "source_attack_ratio")["numberValue"] = ratio
            if any(arg["name"] == "damage_type" for arg in node.get("arguments", [])):
                argument(node, "damage_type")["integerValue"] = 1
            if any(arg["name"] == "reason_param" for arg in node.get("arguments", [])):
                argument(node, "reason_param")["integerValue"] = reason_param
        if trigger_id == 10070212:
            for node in walk(trigger):
                if node.get("type") == "add_buff":
                    value = argument(node, "buff_ids")
                    value["integerListValue"] = [10070201]

    # Sun Shangxiang passive: a basic-attack damage event refreshes skill 1.
    # The runtime currently exposes reset_cooldown as the stable config action;
    # the 500 ms demo parameter remains documented in the level table.
    passive = next(t for t in skills["module"]["triggers"] if t["id"] == 10070011)
    passive["id"] = 10070002
    passive["name"] = "孙尚香-活力迸发-普攻刷新翻滚"
    passive["groupPath"] = "trigger_10070000"
    passive["event"] = "damage.apply.after"
    passive["entryMode"] = "Event"
    passive["scope"] = "owner"
    passive_source_path = TRIGGER_ROOT / "moba-passives.trigger.json"
    passive_source = json.loads(passive_source_path.read_text(encoding="utf-8"))
    passive_source["module"]["triggers"] = deduplicate_trigger_ids(passive_source["module"]["triggers"])
    passive_source["module"]["triggers"] = [
        remap_trigger_namespace(trigger) if 10070000 <= trigger["id"] < 10080000 else trigger
        for trigger in passive_source["module"]["triggers"]
    ]
    condition_source = next(t for t in passive_source["module"]["triggers"] if t["id"] == 10060002)["condition"]
    condition = json.loads(json.dumps(condition_source).replace("1006", "1007"))
    for node in walk(condition):
        if node.get("type") == "arg_eq":
            for arg in node.get("arguments", []):
                if arg["name"] == "right":
                    arg["value"]["integerValue"] = 1
    passive["condition"] = condition
    passive["actions"] = {
        "nodeId": "node_sunshangxiang_reset_cooldown",
        "enabled": True,
        "kind": "Action",
        "type": "reset_cooldown",
        "arguments": [
            {"name": "skill_id", "value": {"source": "Constant", "type": "Integer", "integerValue": 10070101}},
            {"name": "skill_slot", "value": {"source": "Constant", "type": "Integer", "integerValue": 1}},
        ],
        "children": [],
        "elseChildren": [],
    }
    passive["note"] = "孙尚香演示被动：普攻命中刷新翻滚突袭冷却；演示参数 500ms。"
    skills["module"]["triggers"] = deduplicate_trigger_ids(skills["module"]["triggers"])
    skill_path.write_text(json.dumps(skills, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    passive_path = passive_source_path
    passives = passive_source
    for trigger in list(passives["module"]["triggers"]):
        if trigger["id"] not in {10060000, 10060002, 10060003}:
            continue
        clone = json.loads(json.dumps(trigger))
        raw = json.dumps(clone, ensure_ascii=False).replace("1006", "1007")
        clone = json.loads(raw)
        clone["id"] = int(str(trigger["id"]).replace("1006", "1007"))
        clone["name"] = clone["name"].replace("嬴政", "孙尚香")
        clone["groupPath"] = "trigger_10070000"
        if clone["id"] == 10070002:
            # The actual reset action is authored in the skill module; avoid a
            # duplicate passive trigger with the same ID in another package.
            continue
        passives["module"]["triggers"].append(clone)
    passives["module"]["triggers"] = deduplicate_trigger_ids(passives["module"]["triggers"])
    passive_path.write_text(json.dumps(passives, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    append_excel_sources()
    clone_triggers()
    print("added 1007 Sun Shangxiang production source rows and trigger sources")
