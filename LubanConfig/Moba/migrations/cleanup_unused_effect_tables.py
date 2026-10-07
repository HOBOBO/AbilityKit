"""Remove obsolete effect-related sheets from the production workbooks."""

from pathlib import Path

from openpyxl import load_workbook


DATA = Path(__file__).resolve().parents[1] / "Production" / "Datas"
TABLES = {"effects": "效果", "dtbuff": "旧Buff", "ongoing_effects": "持续效果"}


def main():
    workbook_path = DATA / "J-效果表.xlsx"
    workbook = load_workbook(workbook_path)
    removed_sheets = 0
    for sheet_name in TABLES.values():
        if sheet_name in workbook:
            del workbook[sheet_name]
            removed_sheets += 1
    if removed_sheets:
        workbook.save(workbook_path)

    registry_path = DATA / "__tables__.xlsx"
    registry = load_workbook(registry_path)
    sheet = registry.active
    rows = [row for row in range(1, sheet.max_row + 1)
            if sheet.cell(row, 11).value in TABLES]
    for row in reversed(rows):
        sheet.delete_rows(row)
    if rows:
        registry.save(registry_path)

    print(f"removed {removed_sheets} unused effect sheets; {len(rows)} registry entries")


if __name__ == "__main__":
    main()
