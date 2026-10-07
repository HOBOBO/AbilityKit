"""One-time cleanup of duplicated MOBA Excel configuration ownership."""

from copy import copy
from pathlib import Path

from openpyxl import load_workbook


DATA = Path(__file__).resolve().parents[1] / "Production/Datas"


def remove_duplicate_skill_columns():
    path = DATA / "J-角色表.xlsx"
    book = load_workbook(path)
    sheet = book["属性模板"]
    for name in ("PassiveSkills", "ActiveSkills"):
        column = next(cell.column for cell in sheet[1] if cell.value == name)
        sheet.delete_cols(column)
    book.save(path)


def reuse_modifier_definitions():
    path = DATA / "J-效果表.xlsx"
    book = load_workbook(path)
    definitions = book["修饰器定义"]
    headers = {cell.value: cell.column for cell in definitions[1] if cell.value}
    duplicates = {17: 16, 22: 16, 25: 12}
    rows = {definitions.cell(row, headers["Id"]).value: row
            for row in range(4, definitions.max_row + 1)}
    for duplicate, canonical in duplicates.items():
        source = [definitions.cell(rows[duplicate], column).value
                  for column in range(3, definitions.max_column + 1)]
        target = [definitions.cell(rows[canonical], column).value
                  for column in range(3, definitions.max_column + 1)]
        if source != target:
            raise ValueError(f"Modifier {duplicate} no longer matches {canonical}")
    buffs = book["Buff"]
    id_column = next(cell.column for cell in buffs[1] if cell.value == "ModifierIds")
    for row in range(4, buffs.max_row + 1):
        value = buffs.cell(row, id_column).value
        if value:
            ids = [duplicates.get(int(identity), int(identity)) for identity in value.split(",")]
            buffs.cell(row, id_column).value = ",".join(map(str, ids))
    for row in sorted((rows[identity] for identity in duplicates), reverse=True):
        definitions.delete_rows(row)

    tags = book["标签定义"]
    namespace_column = tags.max_column + 1
    tags.cell(1, namespace_column, "Namespace")
    tags.cell(2, namespace_column, "MobaTagNamespace")
    for header_row in (1, 2):
        tags.cell(header_row, namespace_column)._style = copy(tags.cell(header_row, namespace_column - 1)._style)
    tags.column_dimensions[tags.cell(1, namespace_column).column_letter].width = 24
    for row in range(4, tags.max_row + 1):
        identity = tags.cell(row, 2).value
        tags.cell(row, namespace_column, "LegacyNumeric" if identity <= 3 else "GameplayState")
    book.save(path)


def register_skill_tags():
    effect_path = DATA / "J-效果表.xlsx"
    skill_path = DATA / "J-技能表.xlsx"
    effects = load_workbook(effect_path)
    tags = effects["标签定义"]
    skill_book = load_workbook(skill_path, read_only=True)
    skills = skill_book["技能"]
    headers = {cell.value: cell.column for cell in skills[1] if cell.value}
    seen = set()
    for row in range(4, skills.max_row + 1):
        values = skills.cell(row, headers["Tags"]).value
        if not values:
            continue
        for identity in map(int, str(values).split(",")):
            if identity in seen:
                continue
            seen.add(identity)
            destination = tags.max_row + 1
            tags.append([None, identity, skills.cell(row, headers["Name"]).value, "SkillSelection"])
            for column in range(2, 5):
                tags.cell(destination, column)._style = copy(tags.cell(4, min(column, 3))._style)
    effects.save(effect_path)


if __name__ == "__main__":
    remove_duplicate_skill_columns()
    reuse_modifier_definitions()
    register_skill_tags()
