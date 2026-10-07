"""One-time migration of reusable MOBA continuous tag templates."""

from copy import copy
from pathlib import Path

from openpyxl import load_workbook


DATA = Path(__file__).resolve().parents[1] / "Production/Datas"
TAG_BOOK = DATA / "J-标签表.xlsx"
EFFECT_BOOK = DATA / "J-效果表.xlsx"

# Template IDs 3001-3099 belong to shared combat states, independent of heroes.
TEMPLATES = (
    (3001, "通用-控制免疫", "1001", 0, None),
    (3002, "通用-霸体与伤害免疫", "1002,1017", 0, None),
    (3003, "通用-伤害免疫", "1017", 0, None),
    (3004, "通用-不可选取", "1022", 0, None),
    (3005, "通用-普通控制-眩晕", "1020", 10000001, None),
    (3006, "通用-普通控制-禁锢", "1018", 10000001, None),
    (3007, "通用-普通控制-沉默", "1019", 10000001, None),
    (3008, "通用-普通控制-恐惧", "1016", 10000001, None),
    (3009, "通用-普通控制-魅惑", "1014", 10000001, None),
    (3010, "通用-普通控制-睡眠", "1013", 10000001, None),
    (3011, "通用-普通控制-失能", "1015", 10000001, None),
    (3012, "通用-压制标记", "1021", 0, None),
    (3013, "通用-减速标记", "1006", 0, None),
    (3014, "通用-不可选取与伤害免疫", "1017,1022", 0, None),
    (3015, "通用-自主移动执行门禁", None, 0, "5"),
)


def fields(sheet):
    return {cell.value: cell.column for cell in sheet[1] if cell.value and not cell.value.startswith("##")}


def migrate_tags():
    book = load_workbook(TAG_BOOK)
    groups = book["标签规则组"]
    for row in groups.iter_rows(min_row=4):
        if row[1].value == 10000001:
            row[2].value = "通用-目标可受普通控制"

    sheet = book["持续标签模板"]
    columns = fields(sheet)
    existing = {sheet.cell(row, columns["Id"]).value: row for row in range(4, sheet.max_row + 1)}
    sheet.cell(existing[10000001], columns["Name"]).value = "通用-目标可受普通控制的位移"
    sheet.cell(existing[10010001], columns["Name"]).value = "通用-施法期间霸体"
    style_row = existing[10010001]
    for identity, name, tags, activation, blocked in TEMPLATES:
        row = existing.get(identity)
        if row is None:
            row = sheet.max_row + 1
            for column in range(1, sheet.max_column + 1):
                source = sheet.cell(style_row, column)
                target = sheet.cell(row, column)
                target._style = copy(source._style)
                target.number_format = source.number_format
        for field, value in (("Id", identity), ("Name", name), ("ApplicationTagIds", tags),
                             ("ActivationRuleGroupId", activation), ("ExecutionBlockedFlagIds", blocked)):
            sheet.cell(row, columns[field]).value = value
    sheet.auto_filter.ref = f"B4:O{sheet.max_row}"
    book.save(TAG_BOOK)


def migrate_slow_buffs():
    book = load_workbook(EFFECT_BOOK)
    sheet = book["Buff"]
    columns = fields(sheet)
    found = set()
    for row in range(4, sheet.max_row + 1):
        identity = sheet.cell(row, columns["Id"]).value
        if identity not in (10060101, 10060211):
            continue
        if sheet.cell(row, columns["TagIds"]).value not in (None, "1006"):
            raise ValueError(f"Buff/{identity}: unexpected tags before migration")
        sheet.cell(row, columns["TagIds"]).value = None
        sheet.cell(row, columns["ContinuousTagTemplateId"]).value = 3013
        found.add(identity)
    if found != {10060101, 10060211}:
        raise ValueError(f"missing slow buffs: {found}")
    book.save(EFFECT_BOOK)


if __name__ == "__main__":
    migrate_tags()
    migrate_slow_buffs()
