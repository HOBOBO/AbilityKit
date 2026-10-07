"""按逻辑、模型和表现职责整理 MOBA Luban 工作簿。"""

import json
import os
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


ROOT = Path(__file__).resolve().parents[3]
DATAS = ROOT / "LubanConfig/Moba/Production/Datas"
VFX_JSON = ROOT / "Unity/Packages/com.abilitykit.demo.moba.view.runtime/Resources/vfx/vfx.json"
GROUPS = {
    "J-角色表": {
        "characters": "角色",
        "attribute_templates": "属性模板",
        "attr_types": "属性类型",
        "brains": "AI脑",
        "actor_state_machines": "状态机",
        "character_hfsm": "角色HFSM",
        "attributetemplates": "旧属性模板",
    },
    "M-模型表": {
        "models": "模型资源",
        "character_view_actions": "角色动画动作",
    },
    "J-技能表": {
        "skills": "技能",
        "passive_skills": "被动技能",
        "skill_level_tables": "技能等级",
        "search_query_templates": "目标查询",
    },
    "J-技能表现表": {
        "skill_button_templates": "技能按钮与瞄准",
    },
    "T-特效表": {
        "vfx": "特效资源",
        "presentation_templates": "特效表现模板",
    },
    "J-效果表": {
        "buffs": "Buff",
        "effects": "效果",
        "effect_plans": "效果计划",
        "ongoing_effects": "持续效果",
        "continuous_processes": "持续流程",
        "motion_groups": "运动组",
        "component_templates": "组件模板",
        "tag_templates": "标签模板",
        "continuous_tag_templates": "持续标签模板",
        "dtbuff": "旧Buff",
    },
    "J-投射物表": {
        "projectile_launchers": "投射发射器",
        "projectiles": "投射物",
        "aoes": "区域",
        "emitters": "发射器",
    },
    "J-召唤物表": {
        "summons": "召唤物",
        "summon_attr_inherits": "召唤属性继承",
        "spawn_summon_action_templates": "召唤动作模板",
    },
    "J-战斗场景表": {
        "battle_maps": "战斗地图",
        "battle_start": "战斗开局",
        "gameplays": "玩法",
    },
    "J-资源节点表": {
        "ability_nodes": "Ability节点",
        "demo_tbitem": "演示物品",
    },
}


def copy_sheet(source, target):
    for row in source:
        for cell in row:
            destination = target[cell.coordinate]
            destination.value = cell.value
            if cell.has_style:
                destination.font = copy(cell.font)
                destination.fill = copy(cell.fill)
                destination.border = copy(cell.border)
                destination.alignment = copy(cell.alignment)
                destination.number_format = cell.number_format
                destination.protection = copy(cell.protection)
            if cell.hyperlink:
                destination._hyperlink = copy(cell.hyperlink)
            if cell.comment:
                destination.comment = copy(cell.comment)
    for key, dimension in source.row_dimensions.items():
        target.row_dimensions[key] = copy(dimension)
    for key, dimension in source.column_dimensions.items():
        target.column_dimensions[key] = copy(dimension)
    target.freeze_panes = source.freeze_panes
    target.sheet_format = copy(source.sheet_format)
    target.sheet_properties = copy(source.sheet_properties)
    for merged in source.merged_cells.ranges:
        target.merge_cells(str(merged))


def vfx_sheet():
    rows = json.loads(VFX_JSON.read_text(encoding="utf-8-sig"))
    if not rows or any(set(row) != {"Id", "Resource", "DurationMs"} for row in rows):
        raise ValueError("vfx/vfx.json 的字段与迁移模板不一致")
    if len({row["Id"] for row in rows}) != len(rows):
        raise ValueError("vfx/vfx.json 存在重复 Id")
    book = Workbook()
    sheet = book.active
    sheet.title = "vfx"
    sheet.append(["##var", "Id", "Resource", "DurationMs"])
    sheet.append(["##type", "int", "string", "int"])
    sheet.append(["##", "特效资源主源；挂点、跟随和偏移配置见特效表现模板。"])
    for row in rows:
        sheet.append([None, row["Id"], row["Resource"], row["DurationMs"]])
    for cells in sheet:
        for cell in cells:
            cell.font = Font(name="Arial", size=10, bold=cell.row <= 2)
    sheet.freeze_panes = "B4"
    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 48
    sheet.column_dimensions["D"].width = 18
    return sheet


def main():
    catalog_path = DATAS / "__tables__.xlsx"
    catalog = load_workbook(catalog_path)
    entries = {row[10].value: row for row in catalog.active.iter_rows(min_row=4)
               if row[10].value}
    grouped = {name for mapping in GROUPS.values() for name in mapping}
    missing = grouped - set(entries)
    extra = set(entries) - grouped
    if missing not in (set(), {"vfx"}) or extra:
        raise ValueError(f"分组与登记表不一致：{grouped ^ set(entries)}")

    # 先读入所有源表，再覆盖旧工作簿，避免同组迁移时丢失未处理的子表。
    books = {}
    sources = {}
    for name, row in entries.items():
        sheet_name, file_name = row[4].value.split("@", 1)
        if file_name not in books:
            books[file_name] = load_workbook(DATAS / file_name)
        sources[name] = books[file_name][sheet_name]
    if "vfx" not in sources:
        sources["vfx"] = vfx_sheet()

    pending = []
    try:
        for group, mapping in GROUPS.items():
            result = Workbook()
            result.remove(result.active)
            for name, sheet_name in mapping.items():
                copy_sheet(sources[name], result.create_sheet(sheet_name))
            temporary = DATAS / f"~{group}.xlsx"
            result.save(temporary)
            pending.append((temporary, DATAS / f"{group}.xlsx"))
            check = load_workbook(temporary)
            for name, sheet_name in mapping.items():
                original = sources[name]
                target = check[sheet_name]
                if list(original.values) != list(target.values):
                    raise ValueError(f"{name}: 迁移后单元格内容不一致")
                for source_row in original:
                    for cell in source_row:
                        copied = target[cell.coordinate]
                        if (copy(cell.font) != copy(copied.font) or
                                copy(cell.fill) != copy(copied.fill) or
                                copy(cell.alignment) != copy(copied.alignment)):
                            raise ValueError(f"{name}: 迁移后样式不一致：{cell.coordinate}")
            print(f"已校验 {group}.xlsx：{len(mapping)} 个子表")

        if "vfx" not in entries:
            catalog.active.append([None, "Vfx", "DRVfx", True, None, "Id",
                                   None, "c,s,e", None, None, "vfx"])
            entries["vfx"] = catalog.active[catalog.active.max_row]
        for group, mapping in GROUPS.items():
            for name, sheet_name in mapping.items():
                entries[name][4].value = f"{sheet_name}@{group}.xlsx"
        temporary_catalog = DATAS / "~__tables__.xlsx"
        catalog.save(temporary_catalog)
        pending.append((temporary_catalog, catalog_path))
        for temporary, destination in pending:
            os.replace(temporary, destination)
    finally:
        for temporary, _ in pending:
            temporary.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
