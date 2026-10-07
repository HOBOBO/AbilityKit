"""Retire the Excel tag templates after publishing their authoring JSON."""

from pathlib import Path

from openpyxl import load_workbook


DATA = Path(__file__).resolve().parents[1] / "Production" / "Datas"
SHEETS = {"tag_templates": "标签模板", "continuous_tag_templates": "持续标签模板"}


def main():
    workbook_path = DATA / "J-效果表.xlsx"
    workbook = load_workbook(workbook_path)
    removed = 0
    for name in SHEETS.values():
        if name in workbook:
            del workbook[name]
            removed += 1
    if removed:
        workbook.save(workbook_path)

    registry_path = DATA / "__tables__.xlsx"
    registry = load_workbook(registry_path)
    sheet = registry.active
    rows = [index for index in range(1, sheet.max_row + 1)
            if sheet.cell(index, 11).value in SHEETS]
    for index in reversed(rows):
        sheet.delete_rows(index)
    if rows:
        registry.save(registry_path)
    print(f"removed {removed} tag template sheets and {len(rows)} registrations")


if __name__ == "__main__":
    main()
