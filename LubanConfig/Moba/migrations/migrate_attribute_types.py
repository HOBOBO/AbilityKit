"""One-time conversion of the production attribute-type sheet to row definitions."""

from copy import copy
from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path(__file__).resolve().parents[1] / "Production/Datas/J-角色表.xlsx"
ATTRIBUTES = (
    (1, "HP", "当前生命"),
    (2, "MAX_HP", "最大生命"),
    (3, "EXTRA_HP", "额外生命"),
    (4, "PHYSICS_ATTACK", "物理攻击"),
    (5, "MAGIC_ATTACK", "法术攻击"),
    (6, "EXTRA_PHYSICS_ATTACK", "额外物理攻击"),
    (7, "EXTRA_MAGIC_ATTACK", "额外法术攻击"),
    (8, "PHYSICS_DEFENSE", "物理防御"),
    (9, "MAGIC_DEFENSE", "法术防御"),
    (10, "MANA", "当前法力"),
    (11, "MAX_MANA", "最大法力"),
    (12, "CRITICAL_R", "暴击率"),
    (13, "ATTACK_SPEED_R", "攻击速度加成"),
    (14, "COOLDOWN_REDUCE_R", "冷却缩减"),
    (15, "PHYSICS_PENETRATION_R", "物理穿透率"),
    (16, "MAGIC_PENETRATION_R", "法术穿透率"),
    (17, "MOVE_SPEED", "移动速度"),
    (18, "PHYSICS_BLOODSUCKING_R", "物理吸血率"),
    (19, "MAGIC_BLOODSUCKING_R", "法术吸血率"),
    (20, "ATTACK_RANGE", "攻击范围"),
    (21, "PER_SECOND_BLOOD_R", "每秒生命回复率"),
    (22, "PER_SECOND_MANA_R", "每秒法力回复率"),
    (23, "RESILIENCE_R", "韧性"),
)


def main():
    workbook = load_workbook(WORKBOOK)
    sheet = next(sheet for sheet in workbook if sheet.cell(1, 2).value == "Id"
                 and sheet.cell(1, 3).value in ("ActiveSkills", "Key"))
    template = next(sheet for sheet in workbook if
                    "BasicAttackSkillId" in [cell.value for cell in sheet[1]])
    header_style = copy(template.cell(1, 2)._style)
    data_style = copy(template.cell(4, 2)._style)
    old_headers = [cell.value for cell in sheet[1]]
    if "Key" not in old_headers:
        if "Hp" not in old_headers:
            raise ValueError("Unexpected attribute-type source layout")
        sheet.delete_rows(1, sheet.max_row)
        sheet.delete_cols(1, sheet.max_column)
        sheet.append(["##var", "Id", "Key", "Name", "ValueKind", "DefaultValue"])
        sheet.append(["##type", "int", "string", "string", "MobaAttributeValueKind", "float"])
        sheet.append(["##", "ID 与 BattleAttributeType 一致；一个属性一行。", None, None,
                      "Flat 为普通数值，Percent 为比率；计算公式尚由运行时决定。", None])
        for identity, key, name in ATTRIBUTES:
            sheet.append([None, identity, key, name,
                          "Percent" if key.endswith("_R") else "Flat", 0])
    for index, row in enumerate(sheet, 1):
        for cell in row:
            cell._style = copy(header_style if index <= 2 else data_style)
    workbook.save(WORKBOOK)


if __name__ == "__main__":
    main()
