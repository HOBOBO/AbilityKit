"""Add the first formal melee/ranged minion definitions to the MOBA workbooks."""

from copy import copy
from pathlib import Path
import os

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
DATAS = ROOT / "LubanConfig/Moba/Production/Datas"

MINIONS = {
    1: {
        "character_id": 2001,
        "attribute_id": 2001,
        "model_id": 1101,
        "name": "Melee Minion",
        "prefab": "character/character8",
        "scale": 0.72,
        "hp": 1200,
        "attack": 48,
        "physics_defense": 28,
        "magic_defense": 12,
        "move_speed": 3.2,
        "attack_range": 1,
    },
    2: {
        "character_id": 2002,
        "attribute_id": 2002,
        "model_id": 1102,
        "name": "Ranged Minion",
        "prefab": "character/character8",
        "scale": 0.62,
        "hp": 780,
        "attack": 62,
        "physics_defense": 14,
        "magic_defense": 8,
        "move_speed": 3.0,
        "attack_range": 5,
    },
}


def find_sheet(required):
    for path in DATAS.glob("*.xlsx"):
        workbook = load_workbook(path)
        for sheet in workbook.worksheets:
            headers = {cell.value for cell in sheet[1]}
            if required.issubset(headers):
                return path, workbook, sheet
    raise RuntimeError(f"Could not find workbook sheet with columns: {sorted(required)}")


def find_sheet_in_workbook(workbook_name, required):
    path = DATAS / workbook_name
    if not path.is_file():
        raise RuntimeError(f"Could not find workbook: {path}")
    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        headers = {cell.value for cell in sheet[1]}
        if required.issubset(headers):
            return path, workbook, sheet
    raise RuntimeError(
        f"Could not find sheet in {workbook_name} with columns: {sorted(required)}")


def append_row(sheet, values):
    headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
    row = sheet.max_row + 1
    previous = sheet[row - 1]
    for column, cell in enumerate(previous, start=1):
        target = sheet.cell(row=row, column=column)
        if cell.has_style:
            target._style = copy(cell._style)
        if cell.number_format:
            target.number_format = cell.number_format
    for name, value in values.items():
        sheet.cell(row=row, column=headers[name]).value = value


def ensure_rows(sheet, rows):
    headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
    existing = {
        sheet.cell(row=row, column=headers["Id"]).value: row
        for row in range(4, sheet.max_row + 1)
        if sheet.cell(row=row, column=headers["Id"]).value is not None
    }
    for values in rows:
        identity = values["Id"]
        if identity in existing:
            current = existing[identity]
            for name, value in values.items():
                existing_value = sheet.cell(row=current, column=headers[name]).value
                if name == "BasicAttackSkillId" and existing_value == 0:
                    sheet.cell(row=current, column=headers[name]).value = value
                    continue
                if existing_value != value:
                    raise RuntimeError(f"Existing row {identity} has a different {name}")
            continue
        append_row(sheet, values)


def remove_rows_by_id(sheet, ids):
    headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
    id_column = headers["Id"]
    for row in range(sheet.max_row, 3, -1):
        if sheet.cell(row=row, column=id_column).value in ids:
            sheet.delete_rows(row, 1)


def save(workbook, path):
    temporary = path.with_name(f"~{path.name}")
    workbook.save(temporary)
    os.replace(temporary, path)


def main():
    model_path, model_book, model_sheet = find_sheet_in_workbook(
        "M-模型表.xlsx", {"Id", "PrefabPath", "Scale"})
    ensure_rows(model_sheet, [
        {"Id": item["model_id"], "Name": item["name"],
         "PrefabPath": item["prefab"], "Scale": item["scale"]}
        for item in MINIONS.values()
    ])
    save(model_book, model_path)

    character_path, character_book, character_sheet = find_sheet_in_workbook(
        "J-角色表.xlsx",
        {"Id", "AttributeTemplateId", "ModelId", "SkillIds", "PassiveSkillIds"})
    ensure_rows(character_sheet, [
        {"Id": item["character_id"], "AttributeTemplateId": item["attribute_id"],
         "ModelId": item["model_id"], "Name": item["name"],
         "SkillIds": None, "PassiveSkillIds": None}
        for item in MINIONS.values()
    ])
    save(character_book, character_path)

    attribute_path, attribute_book, attribute_sheet = find_sheet_in_workbook(
        "J-角色表.xlsx",
        {"Id", "MaxHp", "PhysicsAttack", "MoveSpeed", "AttackRange"})
    ensure_rows(attribute_sheet, [
        {"Id": item["attribute_id"], "AttackRange": item["attack_range"],
         "AttackSpeedR": 0, "BasicAttackSkillId": 10010001, "CooldownReduceR": 0,
         "CriticalR": 0, "ExtraHp": 0, "ExtraMagicAttack": 0,
         "ExtraPhysicsAttack": 0, "Hp": item["hp"], "MagicAttack": 0,
         "MagicBloodsuckingR": 0, "MagicDefense": item["magic_defense"],
         "MagicPenetrationR": 0, "Mana": 0, "MaxHp": item["hp"],
         "MaxMana": 0, "MoveSpeed": item["move_speed"],
         "PerSecondBloodR": 0, "PerSecondManaR": 0,
         "PhysicsAttack": item["attack"], "PhysicsBloodsuckingR": 0,
         "PhysicsDefense": item["physics_defense"],
         "PhysicsPenetrationR": 0, "ResilienceR": 0}
        for item in MINIONS.values()
    ])
    save(attribute_book, attribute_path)

    # These ids were historically added to the summon workbook by the first
    # draft of this migration. They are character archetypes and must not be
    # published as summon definitions.
    summon_path, summon_book, summon_sheet = find_sheet(
        {"Id", "AttributeTemplateId", "ModelId", "MaxAlivePerOwner", "DespawnOnOwnerDie"})
    remove_rows_by_id(summon_sheet, {item["character_id"] for item in MINIONS.values()})
    save(summon_book, summon_path)


if __name__ == "__main__":
    main()
