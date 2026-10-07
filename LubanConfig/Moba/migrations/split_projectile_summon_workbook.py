"""一次性拆分投射物和召唤物生产工作簿。"""

import os
from pathlib import Path

from openpyxl import load_workbook


DATAS = Path(__file__).resolve().parents[1] / "Production/Datas"
SOURCE = DATAS / "J-投射与召唤表.xlsx"
CATALOG = DATAS / "__tables__.xlsx"
SHEETS = {
    "J-投射物表.xlsx": ("投射发射器", "投射物", "区域", "发射器"),
    "J-召唤物表.xlsx": ("召唤物", "召唤属性继承", "召唤动作模板"),
}


def main():
    if not SOURCE.is_file():
        raise ValueError(f"旧工作簿不存在：{SOURCE}")
    source = load_workbook(SOURCE)
    names = {name for group in SHEETS.values() for name in group}
    if set(source.sheetnames) != names:
        raise ValueError(f"源子表与预期不一致：{set(source.sheetnames) ^ names}")
    catalog = load_workbook(CATALOG)
    registered = {}
    for row in catalog.active.iter_rows(min_row=4):
        value = row[4].value
        if isinstance(value, str) and value.endswith("@" + SOURCE.name):
            sheet_name = value.split("@", 1)[0]
            if sheet_name in registered:
                raise ValueError(f"重复登记子表：{sheet_name}")
            registered[sheet_name] = row[4]
    if set(registered) != names:
        raise ValueError(f"登记表与源子表不一致：{set(registered) ^ names}")

    pending = []
    try:
        for filename, keep in SHEETS.items():
            destination = DATAS / filename
            if destination.exists():
                raise ValueError(f"目标工作簿已存在：{destination}")
            book = load_workbook(SOURCE)
            for sheet in list(book):
                if sheet.title not in keep:
                    book.remove(sheet)
            book._sheets.sort(key=lambda sheet: keep.index(sheet.title))
            temporary = DATAS / ("~" + filename)
            book.save(temporary)
            pending.append((temporary, destination))
            check = load_workbook(temporary)
            if check.sheetnames != list(keep):
                raise ValueError(f"{filename}: 子表顺序不一致")
            for name in keep:
                original, copied = source[name], check[name]
                if (list(original.values) != list(copied.values) or
                        original.freeze_panes != copied.freeze_panes or
                        list(original.merged_cells.ranges) != list(copied.merged_cells.ranges)):
                    raise ValueError(f"{filename}/{name}: 内容或布局不一致")
                for row in original:
                    for cell in row:
                        if cell._style != copied[cell.coordinate]._style:
                            raise ValueError(f"{filename}/{name}: {cell.coordinate} 样式不一致")
                registered[name].value = f"{name}@{filename}"
        temporary_catalog = DATAS / "~__tables__.xlsx"
        catalog.save(temporary_catalog)
        pending.append((temporary_catalog, CATALOG))
        for temporary, destination in pending:
            os.replace(temporary, destination)
    finally:
        for temporary, _ in pending:
            temporary.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
