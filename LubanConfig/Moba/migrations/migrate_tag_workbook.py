"""One-time migration of MOBA tag sources into a dedicated Luban workbook."""

import json
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


ROOT = Path(__file__).resolve().parents[3]
DATAS = ROOT / "LubanConfig/Moba/Production/Datas"
SOURCE = ROOT / "Unity/Packages/com.abilitykit.demo.moba.view.runtime/Configs/MobaTagAuthoring"

ALIASES = {
    "Untargetable": ("State.Untargetable", "Untargetable"),
    "Invulnerable": ("State.Invulnerable", "Invulnerable"),
    "CannotCast": ("State.Stunned", "Stunned", "stunned", "State.Disabled", "Disabled", "disabled",
                   "State.Suppressed", "Suppressed", "suppressed", "State.Silenced", "Silenced", "silenced",
                   "State.Feared", "Feared", "feared", "State.Asleep", "Asleep", "Sleeping", "asleep", "sleeping"),
    "Stunned": ("State.Stunned", "Stunned", "stunned", "State.Disabled", "Disabled", "disabled",
                "State.Suppressed", "Suppressed", "suppressed"),
    "CannotMove": ("State.Stunned", "Stunned", "stunned", "State.Disabled", "Disabled", "disabled",
                   "State.Suppressed", "Suppressed", "suppressed", "State.Rooted", "Rooted", "rooted",
                   "State.Feared", "Feared", "feared", "State.Asleep", "Asleep", "Sleeping", "asleep", "sleeping"),
    "ControlImmune": ("State.ControlImmune", "ControlImmune", "control_immune",
                      "State.SuperArmor", "SuperArmor", "super_armor"),
    "CannotBeControlled": ("State.Stunned", "Stunned", "stunned", "State.Feared", "Feared", "feared",
                           "State.Charmed", "Charmed", "charmed", "State.Asleep", "Asleep", "Sleeping",
                           "asleep", "sleeping"),
}


def add_sheet(book, name, columns, types, comment, rows):
    sheet = book.create_sheet(name)
    sheet.append(["##var", *columns])
    sheet.append(["##type", *types])
    sheet.append(["##", comment])
    for row in rows:
        sheet.append([None, *row])
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Arial", size=10, bold=cell.row <= 2)
    sheet.freeze_panes = "B4"
    sheet.column_dimensions["A"].width = 10
    for col in range(2, len(columns) + 2):
        sheet.column_dimensions[sheet.cell(1, col).column_letter].width = 24


def main():
    if (DATAS / "J-标签表.xlsx").exists():
        raise RuntimeError("Tag workbook already migrated; edit the production workbook directly.")
    old = load_workbook(DATAS / "J-效果表.xlsx")
    definitions = old["标签定义"]
    book = Workbook()
    book.remove(book.active)
    target = book.create_sheet("标签定义")
    for row in definitions:
        for source in row:
            cell = target[source.coordinate]
            cell.value = source.value
            if source.has_style:
                cell._style = copy(source._style)
    for key, dimension in definitions.column_dimensions.items():
        target.column_dimensions[key].width = dimension.width
    target.freeze_panes = "B4"
    old.remove(definitions)

    tags = {row[2]: row[1] for row in list(target.values)[3:] if row[1] is not None}
    next_tag = 1003
    for name in sorted({name for group in ALIASES.values() for name in group} | {"Debuff.Slow"}):
        if name not in tags:
            tags[name] = next_tag
            target.append([None, next_tag, name, "GameplayState"])
            next_tag += 1

    groups, nodes = [], []
    next_node = 1

    def add_group(group_id, name, operator, children):
        nonlocal next_node
        root = next_node
        next_node += 1
        groups.append((group_id, name, root))
        nodes.append((root, group_id, 0, 0, operator, 0, False, 0))
        for order, (child_operator, tag_name, ref_id) in enumerate(children, 1):
            nodes.append((next_node, group_id, root, order, child_operator,
                          tags.get(tag_name, 0), False, ref_id))
            next_node += 1

    for flag_id, (key, names) in enumerate(ALIASES.items(), 1):
        add_group(2000 + flag_id, key + " 标签规则", "Any", [("Tag", name, 0) for name in names])
    add_group(10000001, "控制位移霸体阻断", "Not", [("GroupRef", None, 2006)])

    regular = json.loads((SOURCE / "tag_templates.json").read_text(encoding="utf-8"))
    continuous = json.loads((SOURCE / "continuous_tag_templates.json").read_text(encoding="utf-8"))

    def ids(row, field):
        return ",".join(str(tags[name]) for name in row.get(field, [])) or None

    regular_rows = [(row["Id"], row["Name"], ids(row, "RequiredTags"), ids(row, "BlockedTags"),
                     ids(row, "GrantTags"), ids(row, "RemoveTags"), 0) for row in regular]
    continuous_rows = []
    for row in continuous:
        rule = 10000001 if row["Id"] == 10000001 else 0
        continuous_rows.append((row["Id"], row["Name"],
            ids(row, "ActivationRequiredTags"), ids(row, "ActivationBlockedTags"),
            ids(row, "ApplicationTags"), ids(row, "RemovalRequiredTags"),
            ids(row, "RemovalBlockedTags"), ids(row, "OngoingRequiredTags"),
            ids(row, "OngoingBlockedTags"), ids(row, "RemovalTags"), rule, 0, rule))

    add_sheet(book, "标签规则组", ("Id", "Name", "RootNodeId"), ("int", "string", "int"),
              "Named reusable condition group; nodes are edited below.", groups)
    add_sheet(book, "标签规则节点", ("Id", "GroupId", "ParentId", "Order", "Operator", "TagId", "Exact", "RefGroupId"),
              ("int", "int", "int", "int", "TagRuleOperator", "int", "bool", "int"),
              "Tree rows; Tag references TagId, GroupRef references another group.", nodes)
    add_sheet(book, "结果标志定义", ("Id", "Key", "BitIndex", "Name"),
              ("int", "string", "int", "string"),
              "BitIndex is stable after publication; only results used by code need flags.",
              [(i, key, i - 1, key) for i, key in enumerate(ALIASES, 1)])
    add_sheet(book, "标签结果映射", ("Id", "FlagId", "RuleGroupId"), ("int", "int", "int"),
              "Matching any row for a flag sets that result bit.",
              [(i, i, 2000 + i) for i in range(1, len(ALIASES) + 1)])
    add_sheet(book, "标签模板", ("Id", "Name", "RequiredTagIds", "BlockedTagIds", "GrantTagIds",
                            "RemoveTagIds", "ConditionRuleGroupId"),
              ("int", "string", "(list#sep=,),int", "(list#sep=,),int", "(list#sep=,),int",
               "(list#sep=,),int", "int"),
              "Grant/remove ownership remains source-counted; condition references a rule group.", regular_rows)
    add_sheet(book, "持续标签模板", ("Id", "Name", "ActivationRequiredTagIds", "ActivationBlockedTagIds",
                              "ApplicationTagIds", "RemovalRequiredTagIds", "RemovalBlockedTagIds",
                              "OngoingRequiredTagIds", "OngoingBlockedTagIds", "RemovalTagIds",
                              "ActivationRuleGroupId", "RemovalRuleGroupId", "OngoingRuleGroupId"),
              ("int", "string", *("(list#sep=,),int" for _ in range(8)), "int", "int", "int"),
              "Activation, removal and ongoing each reference a reusable rule group.", continuous_rows)

    registry = load_workbook(DATAS / "__tables__.xlsx")
    sheet = registry.active
    for row in sheet.iter_rows(min_row=4):
        if row[1].value == "TagDefinitions":
            row[4].value = "标签定义@J-标签表.xlsx"
    for class_name, sheet_name, output in (
        ("TagRuleGroups", "标签规则组", "tag_rule_groups"),
        ("TagRuleNodes", "标签规则节点", "tag_rule_nodes"),
        ("TagResultFlags", "结果标志定义", "tag_result_flags"),
        ("TagResultMappings", "标签结果映射", "tag_result_mappings"),
        ("TagTemplates", "标签模板", "tag_templates"),
        ("ContinuousTagTemplates", "持续标签模板", "continuous_tag_templates"),
    ):
        sheet.append([None, class_name, "DR" + class_name, True,
                      sheet_name + "@J-标签表.xlsx", "Id", None, "c,s,e", None, None, output])

    old.save(DATAS / "J-效果表.xlsx")
    book.save(DATAS / "J-标签表.xlsx")
    registry.save(DATAS / "__tables__.xlsx")


if __name__ == "__main__":
    main()
