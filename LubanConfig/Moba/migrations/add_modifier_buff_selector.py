"""Add the Buff template selector to the production modifier definition sheet."""

from copy import copy
from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path(__file__).resolve().parents[1] / "Production/Datas/J-效果表.xlsx"


def main():
    workbook = load_workbook(WORKBOOK)
    sheets = [sheet for sheet in workbook
              if any(cell.value == "SelectorSummonTemplateId" for cell in sheet[1])]
    if len(sheets) != 1:
        raise ValueError("Expected exactly one modifier definition sheet")
    sheet = sheets[0]
    headers = {cell.value: cell.column for cell in sheet[1]}
    if "SelectorBuffTemplateId" in headers:
        return
    source = headers["SelectorSummonTemplateId"]
    target = sheet.max_column + 1
    for row in range(1, sheet.max_row + 1):
        old = sheet.cell(row, source)
        new = sheet.cell(row, target)
        new._style = copy(old._style)
        new.number_format = old.number_format
        if row == 1:
            new.value = "SelectorBuffTemplateId"
        elif row == 2:
            new.value = "int"
        elif row == 3:
            new.value = "0 means any; positive Buff ID matches the BuffApply record"
        elif isinstance(sheet.cell(row, 2).value, int):
            new.value = 0
    workbook.save(WORKBOOK)


if __name__ == "__main__":
    main()
